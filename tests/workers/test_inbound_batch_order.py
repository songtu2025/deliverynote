from tests.support.worker import WorkerCase
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from delivery_note.web.models import (
    Batch,
)

from delivery_note.worker import (
    run_once,
)


class WorkerIntegrationTests(WorkerCase):
    def test_self_operated_multi_file_order_shares_balances_and_exports_all_formats(
        self,
    ):
        rule = self.client.post(
            "/api/self-operated-overreceipt-rule-versions",
            headers=self.headers,
            json={"name": "自营仓批次共享超收 5 件", "allowance": 5},
        )
        self.assertEqual(rule.status_code, 201, rule.text)
        created = self.client.post(
            "/api/self-operated-batches",
            headers=self.headers,
            json={"name": "自营仓多质检单 Worker 测试"},
        )
        self.assertEqual(created.status_code, 201, created.text)
        batch_id = created.json()["id"]

        first_path = self.create_self_operated_delivery(
            self.root / "260817-狂飙-A质检交货单.xlsx",
            10,
        )
        second_path = self.create_self_operated_delivery(
            self.root / "260817-狂飙-B质检交货单.xlsx",
            8,
        )
        uploaded = []
        for source_path in (first_path, second_path):
            with source_path.open("rb") as upload:
                response = self.client.post(
                    f"/api/batches/{batch_id}/files",
                    headers=self.headers,
                    files={"file": (source_path.name, upload)},
                )
            self.assertEqual(response.status_code, 201, response.text)
            uploaded.append(response.json())

        reordered = self.client.put(
            f"/api/batches/{batch_id}/files/order",
            headers=self.headers,
            json={"file_ids": [uploaded[1]["id"], uploaded[0]["id"]]},
        )
        self.assertEqual(reordered.status_code, 200, reordered.text)
        inbound_path = self.create_self_operated_inbound(
            self.root / "自营仓多质检单待入库.xlsx"
        )
        with inbound_path.open("rb") as upload:
            inbound = self.client.post(
                f"/api/self-operated-batches/{batch_id}/inbound-file",
                headers=self.headers,
                files={"file": (inbound_path.name, upload)},
            )
        self.assertEqual(inbound.status_code, 200, inbound.text)

        preflight = self.client.post(
            f"/api/batches/{batch_id}/preflight",
            headers=self.headers,
        )
        self.assertEqual(preflight.status_code, 200, preflight.text)
        compute = self.client.post(
            f"/api/batches/{batch_id}/compute",
            headers=self.headers,
        )
        self.assertEqual(compute.status_code, 202, compute.text)
        self.assertEqual(
            run_once(self.database_url, self.storage_root),
            compute.json()["id"],
        )

        computed = self.client.get(
            f"/api/batches/{batch_id}",
            headers=self.headers,
        ).json()
        self.assertEqual(
            [
                (
                    source["original_name"],
                    source["import_total"],
                    source["manual_total"],
                )
                for source in computed["files"]
            ],
            [
                ("260817-狂飙-B质检交货单.xlsx", 8, 0),
                ("260817-狂飙-A质检交货单.xlsx", 7, 3),
            ],
        )
        self.assertEqual(
            computed["summary"],
            {
                "delivery_total": 18,
                "import_total": 15,
                "manual_total": 3,
                "conserved": True,
            },
        )

        export = self.client.post(
            f"/api/batches/{batch_id}/export",
            headers=self.headers,
        )
        self.assertEqual(export.status_code, 202, export.text)
        self.assertEqual(
            run_once(self.database_url, self.storage_root),
            export.json()["id"],
        )
        self._assert_download_formats(batch_id)

        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            first_export_dir = Path(stored_batch.zip_path).parent
            merged_path = first_export_dir / f"batch-{batch_id}-merged.xlsx"
        merged_path.unlink()
        regenerated = self.client.post(
            f"/api/batches/{batch_id}/export",
            headers=self.headers,
        )
        self.assertEqual(regenerated.status_code, 202, regenerated.text)
        self.assertEqual(regenerated.json()["status"], "queued")
        self.assertEqual(
            run_once(self.database_url, self.storage_root),
            export.json()["id"],
        )
        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            second_export_dir = Path(stored_batch.zip_path).parent
        self.assertNotEqual(second_export_dir, first_export_dir)
        self.assertFalse(first_export_dir.exists())
        self.assertTrue(second_export_dir.is_dir())

    def _assert_download_formats(self, batch_id):
        exported = self.client.get(
            f"/api/batches/{batch_id}",
            headers=self.headers,
        ).json()
        self.assertTrue(exported["download_ready"])
        self.assertTrue(exported["merged_download_ready"])
        self.assertTrue(all(source["download_ready"] for source in exported["files"]))

        merged_response = self.client.get(
            f"/api/batches/{batch_id}/download-merged",
            headers=self.headers,
        )
        self.assertEqual(merged_response.status_code, 200, merged_response.text)
        merged_book = load_workbook(BytesIO(merged_response.content), data_only=True)
        merged_sheet = merged_book["批量入库"]
        self.assertEqual(merged_sheet.max_row, 2)
        self.assertEqual(
            sum(
                merged_sheet.cell(row, 17).value or 0
                for row in range(2, merged_sheet.max_row + 1)
            ),
            15,
        )

        archive_response = self.client.get(
            f"/api/batches/{batch_id}/download",
            headers=self.headers,
        )
        self.assertEqual(archive_response.status_code, 200, archive_response.text)
        with ZipFile(BytesIO(archive_response.content)) as archive:
            self.assertEqual(
                sorted(archive.namelist()),
                [
                    "260817-狂飙-A质检交货单_积加入库.xlsx",
                    "260817-狂飙-B质检交货单_积加入库.xlsx",
                ],
            )

        for source in exported["files"]:
            response = self.client.get(
                f"/api/batch-files/{source['id']}/download",
                headers=self.headers,
            )
            self.assertEqual(response.status_code, 200, response.text)
