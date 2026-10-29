from tests.support.worker import WorkerCase
from io import BytesIO

from openpyxl import load_workbook
from tests.support.inbound_inputs import create_ambiguous_product

from delivery_note.web.models import (
    Batch,
    BatchFile,
)

from delivery_note.worker import (
    run_once,
)


class WorkerIntegrationTests(WorkerCase):
    def test_self_operated_site_choice_recomputes_and_exports_inbound_workbook(self):
        self._prepare_site_choice_inputs()

        created = self.client.post(
            "/api/self-operated-batches",
            headers=self.headers,
            json={"name": "自营仓 Worker 集成测试"},
        )
        self.assertEqual(created.status_code, 201, created.text)
        batch_id = created.json()["id"]
        delivery = self.create_self_operated_delivery(
            self.root / "260817-狂飙-质检交货单.xlsx",
            18,
        )
        inbound = self.create_self_operated_inbound(self.root / "自营仓收货入库单.xlsx")
        with delivery.open("rb") as upload:
            source = self.client.post(
                f"/api/batches/{batch_id}/files",
                headers=self.headers,
                files={"file": (delivery.name, upload)},
            )
        self.assertEqual(source.status_code, 201, source.text)
        with inbound.open("rb") as upload:
            inbound_upload = self.client.post(
                f"/api/self-operated-batches/{batch_id}/inbound-file",
                headers=self.headers,
                files={"file": (inbound.name, upload)},
            )
        self.assertEqual(inbound_upload.status_code, 200, inbound_upload.text)
        preflight = self.client.post(
            f"/api/batches/{batch_id}/preflight",
            headers=self.headers,
        )
        self.assertEqual(preflight.status_code, 200, preflight.text)
        run_once(self.database_url, self.storage_root)
        rows = self.client.get(
            f"/api/batches/{batch_id}/receipts", headers=self.headers
        ).json()["items"]
        self.assertEqual(len(rows), 1)
        self.assertIn("站点不唯一", rows[0]["issue"])
        self.assertIsNone(rows[0]["requested"])
        blocked = self.client.post(
            f"/api/batches/{batch_id}/compute", headers=self.headers
        )
        self.assertEqual(blocked.status_code, 409)
        endpoint = f"/api/batches/{batch_id}/files/{rows[0]['file_id']}/receipt-site"
        payload = {
            "sku": "SKU-A",
            "original_site": "US",
            "full_site": "AMAZON:OTHER:US",
        }
        rejected = self.client.put(endpoint, headers=self.headers, json=payload)
        self.assertEqual(rejected.status_code, 400)
        payload["full_site"] = "AMAZON:RIVMOUNT:US"
        selected = self.client.put(endpoint, headers=self.headers, json=payload)
        self.assertEqual(selected.status_code, 200, selected.text)
        run_once(self.database_url, self.storage_root)
        with self.app.state.database.session() as session:
            source = session.get(BatchFile, rows[0]["file_id"])
            self.assertEqual(source.receipt_sites, {"SKU-A\tUS": "AMAZON:RIVMOUNT:US"})
        compute = self.client.post(
            f"/api/batches/{batch_id}/compute", headers=self.headers
        )
        self.assertEqual(compute.status_code, 202, compute.text)
        run_once(self.database_url, self.storage_root)

        batch = self.client.get(
            f"/api/batches/{batch_id}",
            headers=self.headers,
        ).json()
        self.assertEqual(batch["summary"]["delivery_total"], 18)
        self.assertEqual(batch["summary"]["import_total"], 15)
        self.assertEqual(batch["summary"]["manual_total"], 3)
        exceptions = self.client.get(
            f"/api/batches/{batch_id}/exceptions",
            headers=self.headers,
        ).json()
        self.assertEqual(len(exceptions), 1)
        self.assertEqual(exceptions[0]["reason"], "超出允许超收量")
        self.assertEqual(exceptions[0]["manual_quantity"], 3)

        self._assert_inbound_export(batch_id, batch)

    def _prepare_site_choice_inputs(
        self,
    ):
        product = create_ambiguous_product(self.root / "product-ambiguous.xlsx")
        with product.open("rb") as upload:
            product_version = self.client.post(
                "/api/input-versions/product",
                headers=self.headers,
                data={"name": "product-ambiguous", "activate": "true"},
                files={"file": (product.name, upload)},
            )
        self.assertEqual(product_version.status_code, 201, product_version.text)
        rule = self.client.post(
            "/api/self-operated-overreceipt-rule-versions",
            headers=self.headers,
            json={"name": "自营仓超收 5 件", "allowance": 5},
        )
        self.assertEqual(rule.status_code, 201, rule.text)

    def _assert_inbound_export(self, batch_id, batch):
        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            stored_batch.error_message = "历史导出失败"
            session.commit()
        export = self.client.post(
            f"/api/batches/{batch_id}/export",
            headers=self.headers,
        )
        self.assertEqual(export.status_code, 202, export.text)
        self.assertEqual(
            run_once(self.database_url, self.storage_root),
            export.json()["id"],
        )
        refreshed_batch = self.client.get(
            f"/api/batches/{batch_id}", headers=self.headers
        ).json()
        self.assertIsNone(refreshed_batch["error_message"])
        result = self.client.get(
            f"/api/batch-files/{batch['files'][0]['id']}/download",
            headers=self.headers,
        )
        self.assertEqual(result.status_code, 200, result.text)
        output = load_workbook(BytesIO(result.content), data_only=True)
        output_sheet = output["批量入库"]
        self.assertEqual(output_sheet.cell(2, 8).value, "AMAZON:RIVMOUNT:US")
        self.assertEqual(output_sheet.cell(2, 12).value, 10)
        self.assertEqual(output_sheet.cell(2, 17).value, 15)
        self.assertEqual(output_sheet.cell(2, 18).value, "未分配库位")
        self.assertIsNone(output_sheet.cell(2, 19).value)
        self.assertEqual(output_sheet.cell(2, 20).value, "规则允许超收：5")
