from tests.support.worker import WorkerCase
from io import BytesIO
from mimetypes import guess_type
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

import delivery_note.workers.export_files as export_files_module
from delivery_note.web.models import (
    Batch,
    ExceptionRecord,
    Job,
)

from delivery_note.worker import (
    run_once,
)


class WorkerIntegrationTests(WorkerCase):
    def test_compute_split_export_and_zip_are_end_to_end(self):
        first = self.create_delivery(
            self.root / "260717-狂飙-A交货单-发货10箱.xlsx", 80
        )
        second = self.create_delivery(
            self.root / "260717-狂飙-B交货单-发货20箱.xlsx", 80
        )
        batch_id, compute_job_id = self.create_batch([first, second])

        self.assertEqual(run_once(self.database_url, self.storage_root), compute_job_id)
        batch = self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()
        self.assertEqual(batch["status"], "succeeded")
        illegal_preflight = self.client.post(
            f"/api/batches/{batch_id}/preflight", headers=self.headers
        )
        self.assertEqual(illegal_preflight.status_code, 409)

        self.assertEqual(
            self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()[
                "status"
            ],
            "succeeded",
        )
        self.assertEqual(
            [(item["import_total"], item["manual_total"]) for item in batch["files"]],
            [(80, 0), (20, 60)],
        )
        exceptions = self.client.get(
            f"/api/batches/{batch_id}/exceptions", headers=self.headers
        ).json()
        self.assertEqual(len(exceptions), 1)
        self.assertEqual(exceptions[0]["manual_quantity"], 60)
        self.assertEqual(exceptions[0]["reason_code"], "purchase_balance_exceeded")
        with self.app.state.database.session() as session:
            stored = session.get(ExceptionRecord, exceptions[0]["id"])
            self.assertEqual(stored.reason_code, "purchase_balance_exceeded")
        self.assertEqual(exceptions[0]["scale_position"], "短尾")
        self.assertEqual(exceptions[0]["stocking_position"], "备货")
        self.assertNotIn("ordered_days", exceptions[0])

        split = self.client.put(
            f"/api/exceptions/{exceptions[0]['id']}/split",
            headers=self.headers,
            json={
                "parts": [
                    {
                        "quantity": 25,
                        "destination": "水鞋-广州仓",
                        "delivery_note": "超出采购未交量",
                        "resolved": True,
                    },
                    {
                        "quantity": 35,
                        "delivery_note": "超出采购未交量",
                        "resolved": False,
                    },
                ]
            },
        )
        self.assertEqual(split.status_code, 200, split.text)
        self.assertEqual(split.json()["scale_position"], "短尾")
        self.assertEqual(split.json()["stocking_position"], "备货")
        self.assertNotIn("ordered_days", split.json())
        export = self.client.post(
            f"/api/batches/{batch_id}/export", headers=self.headers
        )
        self.assertEqual(export.status_code, 202, export.text)
        export_job_id = export.json()["id"]
        self.assertEqual(run_once(self.database_url, self.storage_root), export_job_id)

        self._assert_download_formats(batch_id, export_job_id)

        self._assert_export_regeneration(batch_id, export_job_id)

    def _assert_download_formats(self, batch_id, export_job_id):
        completed = self.client.get(
            f"/api/batches/{batch_id}", headers=self.headers
        ).json()
        self.assertTrue(completed["download_ready"])
        self.assertTrue(completed["merged_download_ready"])
        self.assertTrue(all(item["download_ready"] for item in completed["files"]))
        merged_response = self.client.get(
            f"/api/batches/{batch_id}/download-merged", headers=self.headers
        )
        self.assertEqual(merged_response.status_code, 200, merged_response.text)
        self.assertEqual(
            merged_response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            f"batch-{batch_id}-merged.xlsx",
            merged_response.headers["content-disposition"],
        )
        merged_book = load_workbook(
            BytesIO(merged_response.content),
            data_only=True,
        )
        merged_import_sheet = merged_book["交货导入"]
        merged_pending_sheet = merged_book["待处理导入"]
        self.assertEqual(merged_import_sheet.cell(3, 1).value, "示例仓")
        self.assertEqual(
            sum(
                merged_import_sheet.cell(row, 4).value or 0
                for row in range(4, merged_import_sheet.max_row + 1)
            ),
            125,
        )
        self.assertEqual(
            sum(
                merged_pending_sheet.cell(row, 4).value or 0
                for row in range(3, merged_pending_sheet.max_row + 1)
            ),
            35,
        )
        merged_import_notes = [
            merged_import_sheet.cell(row, 6).value
            for row in range(4, merged_import_sheet.max_row + 1)
        ]
        self.assertTrue(merged_import_notes[0].endswith("-01-10箱"))
        self.assertTrue(
            all(note.endswith("-02-20箱") for note in merged_import_notes[1:])
        )
        self.assertTrue(merged_pending_sheet.cell(3, 6).value.endswith("-02-20箱"))
        archive_response = self.client.get(
            f"/api/batches/{batch_id}/download", headers=self.headers
        )
        self.assertEqual(archive_response.status_code, 200, archive_response.text)
        self.assertEqual(
            archive_response.headers["content-type"], guess_type("export.zip")[0]
        )
        self.assertTrue(
            archive_response.headers["content-disposition"].startswith("attachment;")
        )
        with ZipFile(BytesIO(archive_response.content)) as archive:
            names = sorted(archive.namelist())
            self.assertEqual(len(names), 2)
            self.assertFalse(any("merged" in name for name in names))
            second_book = load_workbook(BytesIO(archive.read(names[1])), data_only=True)

        import_sheet = second_book["交货导入"]
        pending_sheet = second_book["待处理导入"]
        import_total = sum(
            import_sheet.cell(row, 4).value or 0
            for row in range(4, import_sheet.max_row + 1)
        )
        pending_total = sum(
            pending_sheet.cell(row, 4).value or 0
            for row in range(3, pending_sheet.max_row + 1)
        )
        self.assertEqual((import_total, pending_total), (45, 35))
        import_records = [
            [import_sheet.cell(row, column).value for column in range(1, 8)]
            for row in range(4, import_sheet.max_row + 1)
        ]
        self.assertEqual(len(import_records), 1)
        self.assertEqual(import_records[0][3], 45)
        self.assertEqual(import_records[0][6], "超出采购未交量：60")
        self.assertIsNone(run_once(self.database_url, self.storage_root))
        repeated = self.client.post(
            f"/api/batches/{batch_id}/export", headers=self.headers
        )
        self.assertEqual(repeated.json()["id"], export_job_id)
        self.assertEqual(repeated.json()["status"], "succeeded")

    def _assert_export_regeneration(self, batch_id, export_job_id):
        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            first_export_dir = Path(stored_batch.zip_path).parent
            merged_path = Path(stored_batch.zip_path).with_name(
                f"batch-{batch_id}-merged.xlsx"
            )
        merged_path.unlink()
        self.assertEqual(
            self.client.get(
                f"/api/batches/{batch_id}/download-merged",
                headers=self.headers,
            ).status_code,
            404,
        )
        regenerated = self.client.post(
            f"/api/batches/{batch_id}/export", headers=self.headers
        )
        self.assertEqual(regenerated.json()["id"], export_job_id)
        self.assertEqual(regenerated.json()["status"], "queued")
        self.assertEqual(run_once(self.database_url, self.storage_root), export_job_id)
        self.assertEqual(
            self.client.get(
                f"/api/batches/{batch_id}/download-merged",
                headers=self.headers,
            ).status_code,
            200,
        )
        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            second_export_dir = Path(stored_batch.zip_path).parent
        self.assertNotEqual(second_export_dir, first_export_dir)
        self.assertFalse(first_export_dir.exists())
        self.assertTrue(second_export_dir.is_dir())

        second_merged_path = second_export_dir / f"batch-{batch_id}-merged.xlsx"
        second_merged_path.unlink()
        retried = self.client.post(
            f"/api/batches/{batch_id}/export", headers=self.headers
        )
        self.assertEqual(retried.json()["status"], "queued")
        with (
            patch.object(
                export_files_module.shutil,
                "rmtree",
                side_effect=PermissionError("denied"),
            ),
            self.assertLogs("delivery_note.worker", level="WARNING"),
        ):
            self.assertEqual(
                run_once(self.database_url, self.storage_root),
                export_job_id,
            )
        with self.app.state.database.session() as session:
            stored_batch = session.get(Batch, batch_id)
            stored_job = session.get(Job, export_job_id)
            third_export_dir = Path(stored_batch.zip_path).parent
            self.assertEqual(stored_batch.status, "succeeded")
            self.assertEqual(stored_job.status, "succeeded")
        self.assertNotEqual(third_export_dir, second_export_dir)
        self.assertTrue(third_export_dir.is_dir())
        self.assertTrue(second_export_dir.is_dir())
