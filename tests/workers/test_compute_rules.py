from tests.support.worker import WorkerCase
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from delivery_note.processing.models import IMPORT_COLUMNS
from delivery_note.web.models import (
    ExceptionRecord,
)

from delivery_note.worker import (
    run_once,
)


class WorkerIntegrationTests(WorkerCase):
    def test_worker_compute_uses_supplier_alias_from_locked_version(self):
        delivery = self.create_delivery(
            self.root / "260903-瑞智雅RIVBOS眼镜交货单-发货53箱 (1).xlsx",
            40,
        )
        batch_id, compute_job_id = self.create_batch([delivery])

        self.assertEqual(run_once(self.database_url, self.storage_root), compute_job_id)
        batch = self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()

        self.assertEqual(batch["status"], "succeeded")
        self.assertEqual(batch["files"][0]["supplier_code"], "GYS-023")
        self.assertEqual(batch["files"][0]["supplier_name"], "KuangBiao")

    def test_compute_uses_the_overreceipt_rule_locked_when_batch_was_created(self):
        first_rule = self.client.post(
            "/api/overreceipt-rule-versions",
            headers=self.headers,
            json={
                "name": "允许短尾超收 50",
                "short_tail_limit": 50,
                "medium_tail_limit": 20,
                "long_tail_limit": 10,
                "allowed_warehouses": ["水鞋-广州仓"],
            },
        )
        self.assertEqual(first_rule.status_code, 201, first_rule.text)
        first = self.create_delivery(
            self.root / "260717-狂飙-A交货单-发货10箱.xlsx", 80
        )
        second = self.create_delivery(
            self.root / "260717-狂飙-B交货单-发货20箱.xlsx", 80
        )
        batch_id, compute_job_id = self.create_batch([first, second])

        replacement_rule = self.client.post(
            "/api/overreceipt-rule-versions",
            headers=self.headers,
            json={
                "name": "停止自动超收",
                "short_tail_limit": 0,
                "medium_tail_limit": 0,
                "long_tail_limit": 0,
                "allowed_warehouses": [],
            },
        )
        self.assertEqual(replacement_rule.status_code, 201, replacement_rule.text)

        self.assertEqual(run_once(self.database_url, self.storage_root), compute_job_id)
        batch = self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()

        self.assertEqual(batch["overreceipt_rule"]["id"], first_rule.json()["id"])
        self.assertEqual(
            [(item["import_total"], item["manual_total"]) for item in batch["files"]],
            [(80, 0), (70, 10)],
        )
        self.assertEqual(
            batch["files"][1]["import_total"] + batch["files"][1]["manual_total"],
            80,
        )
        exceptions = self.client.get(
            f"/api/batches/{batch_id}/exceptions", headers=self.headers
        ).json()
        self.assertEqual(exceptions[0]["reason"], "超出允许超收量")
        with self.app.state.database.session() as session:
            stored = session.get(ExceptionRecord, exceptions[0]["id"])
            self.assertEqual(stored.reason_code, "overreceipt_limit_exceeded")
        self.assertEqual(exceptions[0]["manual_quantity"], 10)
        self.assertEqual(exceptions[0]["purchase_allocated_quantity"], 20)
        self.assertEqual(exceptions[0]["overreceipt_allocated_quantity"], 50)
        self.assertEqual(exceptions[0]["overreceipt_remaining_quantity"], 0)

        export = self.client.post(
            f"/api/batches/{batch_id}/export", headers=self.headers
        )
        self.assertEqual(export.status_code, 202, export.text)
        self.assertEqual(
            run_once(self.database_url, self.storage_root),
            export.json()["id"],
        )
        archive_response = self.client.get(
            f"/api/batches/{batch_id}/download", headers=self.headers
        )
        with ZipFile(BytesIO(archive_response.content)) as archive:
            names = sorted(archive.namelist())
            second_book = load_workbook(
                BytesIO(archive.read(names[1])),
                data_only=True,
            )
        import_sheet = second_book["交货导入"]
        pending_sheet = second_book["待处理导入"]
        self.assertEqual(
            [import_sheet.cell(2, column).value for column in range(1, 8)],
            IMPORT_COLUMNS,
        )
        self.assertEqual(
            sum(
                import_sheet.cell(row, 4).value or 0
                for row in range(4, import_sheet.max_row + 1)
            ),
            70,
        )
        self.assertEqual(pending_sheet.cell(3, 4).value, 10)
        self.assertEqual(pending_sheet.cell(3, 7).value, "超出允许超收量：10")
