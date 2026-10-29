"""累计实收经过真实路由、Worker 和 Excel 导出的集成验收。"""

from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import select

from delivery_note.web.models import BatchFile, InputVersion, SelfOperatedBatch
from delivery_note.worker import run_once
from tests.inbound.test_history import inbound_order
from tests.support.worker import WorkerCase


class ReceiptWorkerTests(WorkerCase):
    def test_unmatched_receipt_can_compute_and_export_without_importing_it(self):
        batch_id = self.create_receipt_batch(managed=True)
        order = inbound_order()
        order.update(supplierName="KuangBiao", releatedCode="LN2608179025")
        order["orderItemResultList"][0]["sku"] = "OTHER-SKU"
        with patch(
            "delivery_note.workers.prepare_inbound.GerpgoClient.from_config"
        ) as factory:
            factory.return_value.list_self_operated_inbound_orders.return_value = [
                order
            ]
            batch = self.preflight(batch_id)
        self.assertEqual(batch["status"], "preflight_ready")
        self.assertEqual(batch["receipt"]["requested_total"], 19)
        for kind in ("compute", "export"):
            queued = self.client.post(
                f"/api/batches/{batch_id}/{kind}", headers=self.headers
            )
            self.assertEqual(queued.status_code, 202, queued.text)
            run_once(self.database_url, self.storage_root)
        batch = self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()
        self.assertEqual(batch["summary"]["import_total"], 0)
        self.assertEqual(batch["summary"]["manual_total"], 19)
        pending = self.client.get(
            f"/api/batches/{batch_id}/exceptions", headers=self.headers
        ).json()
        self.assertEqual(pending[0]["manual_quantity"], 19)
        self.assertEqual(pending[0]["overreceipt_allocated_quantity"], 0)
        result = self.client.get(
            f"/api/batches/{batch_id}/download", headers=self.headers
        )
        self.assertEqual(result.status_code, 200)
        book = load_workbook(BytesIO(result.content), data_only=True)
        self.assertEqual(book.worksheets[0].max_row, 1)
        book.close()

    def test_missing_source_fails_preparation_and_blocks_compute_and_download(self):
        batch_id = self.create_receipt_batch()
        with self.app.state.database.session() as session:
            profile = session.get(SelfOperatedBatch, batch_id)
            Path(profile.inbound_storage_path).unlink()
        with self.assertLogs("delivery_note.worker", level="ERROR"):
            batch = self.preflight(batch_id)
        self.assertEqual(batch["status"], "failed")
        self.assertEqual(batch["jobs"]["prepare"]["status"], "failed")
        self.assertNotIn("compute", batch["jobs"])
        self.assertIn("自营仓收货入库单不存在", batch["error_message"])
        self.assertFalse(batch["download_ready"])
        for method, endpoint in (("POST", "compute"), ("GET", "download")):
            response = self.client.request(
                method, f"/api/batches/{batch_id}/{endpoint}", headers=self.headers
            )
            self.assertEqual(response.status_code, 409, response.text)

    def test_zero_new_quantity_exports_headers_without_a_fake_product_row(self):
        batch_id = self.create_receipt_batch(cumulative=3, history=3)
        batch = self.preflight(batch_id)
        self.assertEqual(batch["receipt"]["requested_total"], 0)
        for kind in ("compute", "export"):
            queued = self.client.post(
                f"/api/batches/{batch_id}/{kind}", headers=self.headers
            )
            self.assertEqual(queued.status_code, 202, queued.text)
            run_once(self.database_url, self.storage_root)
        result = self.client.get(
            f"/api/batches/{batch_id}/download", headers=self.headers
        )
        book = load_workbook(BytesIO(result.content), data_only=True)
        self.assertEqual(book.worksheets[0].max_row, 1)
        book.close()

    def create_receipt_batch(self, cumulative=19, history=3, *, managed=False):
        delivery = self.create_self_operated_delivery(
            self.root / "狂飙-累计.xlsx", cumulative
        )
        inbound = self.create_self_operated_inbound(self.root / "完整历史.xlsx")
        book = load_workbook(inbound)
        sheet = book.worksheets[0]
        columns = [cell.value for cell in sheet[1]]
        for row in (2, 3):
            for column, value in (
                ("已入库", history),
                ("原始应收货", 19),
                ("应收货", 19 - history),
            ):
                sheet.cell(row, columns.index(column) + 1, value)
        book.save(inbound)
        book.close()
        with delivery.open("rb") as upload, inbound.open("rb") as source:
            result = self.client.post(
                "/api/self-operated-batches",
                headers=self.headers,
                data={"name": "累计实收验收"},
                files={
                    "delivery_file": (delivery.name, upload),
                    "inbound_file": (inbound.name, source),
                },
            )
        self.assertEqual(result.status_code, 201, result.text)
        batch_id = result.json()["id"]
        if managed:
            with self.app.state.database.session() as session:
                profile = session.get(SelfOperatedBatch, batch_id)
                session.add(
                    InputVersion(
                        kind="self_operated_inbound",
                        name="API来源",
                        original_name="API来源.xlsx",
                        storage_path=profile.inbound_storage_path,
                        created_by=result.json()["created_by"],
                        active=True,
                    )
                )
                session.commit()
        return batch_id

    def preflight(self, batch_id):
        result = self.client.post(
            f"/api/batches/{batch_id}/preflight", headers=self.headers
        )
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["status"], "queued")
        run_once(self.database_url, self.storage_root)
        return self.client.get(f"/api/batches/{batch_id}", headers=self.headers).json()

    def test_only_new_quantity_is_exported_and_refresh_invalidates_downloads(self):
        batch_id = self.create_receipt_batch()
        batch = self.preflight(batch_id)
        self.assertEqual(batch["status"], "preflight_ready")
        self.assertEqual(
            (
                batch["receipt"]["cumulative_total"],
                batch["receipt"]["history_total"],
                batch["receipt"]["requested_total"],
            ),
            (19, 3, 16),
        )
        for kind in ("compute", "export"):
            queued = self.client.post(
                f"/api/batches/{batch_id}/{kind}", headers=self.headers
            )
            self.assertEqual(queued.status_code, 202, queued.text)
            run_once(self.database_url, self.storage_root)
        result = self.client.get(
            f"/api/batches/{batch_id}/download", headers=self.headers
        )
        self.assertEqual(result.status_code, 200)
        book = load_workbook(BytesIO(result.content), data_only=True)
        self.assertEqual(book.worksheets[0].cell(2, 17).value, 16)
        book.close()
        refreshed = self.client.post(
            f"/api/batches/{batch_id}/receipts/refresh", headers=self.headers
        )
        self.assertEqual(refreshed.status_code, 200, refreshed.text)
        self.assertEqual(
            self.client.get(
                f"/api/batches/{batch_id}/download", headers=self.headers
            ).status_code,
            409,
        )

    def test_api_snapshot_is_reused_and_ignores_old_global_balance(self):
        batch_id = self.create_receipt_batch(managed=True)
        order = inbound_order()
        order.update(supplierName="KuangBiao", releatedCode="LN2608179025")
        client_path = "delivery_note.workers.prepare_inbound.GerpgoClient.from_config"
        with patch(client_path) as factory:
            factory.return_value.list_self_operated_inbound_orders.return_value = [
                order
            ]
            batch = self.preflight(batch_id)
            self.assertEqual(batch["receipt"]["requested_total"], 16)
            repeated = self.preflight(batch_id)
            self.assertEqual(
                repeated["receipt"]["captured_at"], batch["receipt"]["captured_at"]
            )
            factory.return_value.list_self_operated_inbound_orders.assert_called_once_with(
                bill_codes=["LN2608179025"]
            )
        rows = self.client.get(
            f"/api/batches/{batch_id}/receipts?page_size=1", headers=self.headers
        ).json()
        self.assertEqual(rows["total"], 1)
        self.assertEqual(rows["items"][0]["history"], 3)

    def test_unknown_history_never_creates_pending_quantity(self):
        batch_id = self.create_receipt_batch(history=3, cumulative=2)
        batch = self.preflight(batch_id)
        self.assertEqual(batch["status"], "draft")
        self.assertIsNone(batch["receipt"]["requested_total"])
        self.assertEqual(batch["receipt"]["issue_count"], 1)
        self.assertEqual(
            self.client.post(
                f"/api/batches/{batch_id}/compute", headers=self.headers
            ).status_code,
            409,
        )
        with self.app.state.database.session() as session:
            source = session.scalar(
                select(BatchFile).where(BatchFile.batch_id == batch_id)
            )
            self.assertEqual((source.manual_total, source.import_rows), (0, []))
            self.assertIsNone(source.receipt_context["rows"][0]["requested"])

    def test_old_manual_source_cannot_silently_assume_zero_history(self):
        batch_id = self.create_receipt_batch()
        with self.app.state.database.session() as session:
            profile = session.get(SelfOperatedBatch, batch_id)
            path = Path(profile.inbound_storage_path)
        book = load_workbook(path)
        sheet = book.worksheets[0]
        columns = [cell.value for cell in sheet[1]]
        sheet.delete_cols(columns.index("积加明细ID") + 1)
        book.save(path)
        book.close()
        with self.assertLogs("delivery_note.worker", level="ERROR"):
            batch = self.preflight(batch_id)
        self.assertEqual(batch["status"], "failed")
        self.assertIn("重新取数", batch["error_message"])
        self.assertIsNone(batch["receipt"])
