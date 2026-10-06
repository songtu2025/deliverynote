"""验证只读准备、历史内容核对，以及单文件发布和失败保护。"""

import os
from pathlib import Path
import subprocess
import sys
from unittest.mock import patch

from openpyxl import load_workbook

from delivery_note.web.models import Batch, Job
from delivery_note.workers.export_delivery import _write_merged_delivery
from scripts.backup.archive import sha256
from scripts.merged_export_recovery import prepare_export, publish_export
from tests.support.delivery_exports import assert_delivery_exports
from tests.support.merged_recovery import MergedRecoveryCase


class MergedExportRecoveryTests(MergedRecoveryCase):
    def prepare(self) -> dict:
        return prepare_export(
            self.database_url, self.storage_root, self.batch_id, self.candidate
        )

    def publish(self) -> Path:
        return publish_export(self.database_url, self.storage_root, self.candidate)

    def test_prepare_preserves_database_and_all_historical_files(self) -> None:
        before = self.storage_snapshot()
        report = self.prepare()
        self.assertEqual(self.storage_snapshot(), before)
        self.assertFalse(self.target.exists())
        self.assertEqual(
            (report["delivery_total"], report["import_total"], report["pending_total"]),
            (160, 125, 35),
        )
        self.assertEqual(report["legacy_layout_count"], 0)
        assert_delivery_exports(
            self,
            (self.candidate / self.target.name).read_bytes(),
            self.archive.read_bytes(),
        )
        second = prepare_export(
            self.database_url,
            self.storage_root,
            self.batch_id,
            self.root / "second-candidate",
        )
        self.assertEqual(second["state_sha256"], report["state_sha256"])
        self.assertEqual(self.storage_snapshot(), before)

    def test_legacy_rows_are_compared_without_dropping_first_delivery(self) -> None:
        self.use_legacy_layout()
        before = self.storage_snapshot()
        report = self.prepare()
        self.assertEqual(report["legacy_layout_count"], 2)
        self.assertEqual(report["import_total"], 125)
        self.assertEqual(report["pending_total"], 35)
        self.assertEqual(self.storage_snapshot(), before)

    def test_publish_adds_only_merged_file_and_preserves_downloads_and_jobs(
        self,
    ) -> None:
        before = self.storage_snapshot()
        self.prepare()
        expected_hash = sha256(self.candidate / self.target.name)
        self.assertEqual(self.publish(), self.target)
        after = self.storage_snapshot()
        self.assertEqual(
            after.pop(self.target.relative_to(self.root).as_posix()),
            expected_hash,
        )
        self.assertEqual(after, before)
        response = self.client.get(
            f"/api/batches/{self.batch_id}/download-merged", headers=self.headers
        )
        self.assertEqual(response.status_code, 200)
        assert_delivery_exports(self, response.content, self.archive.read_bytes())
        self.assertFalse(list(self.target.parent.glob(".merged-recovery-*")))
        with self.assertRaises(ValueError):
            self.publish()
        self.assertEqual(
            self.storage_snapshot(),
            {
                **before,
                self.target.relative_to(self.root).as_posix(): expected_hash,
            },
        )

    def test_existing_target_and_storage_candidate_directory_are_refused(self) -> None:
        for directory in (self.storage_root / "candidate", self.root):
            with (
                self.subTest(directory=directory),
                self.assertRaises((ValueError, FileExistsError)),
            ):
                prepare_export(
                    self.database_url, self.storage_root, self.batch_id, directory
                )
        self.target.write_bytes(b"already-present")
        before = self.storage_snapshot()
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertEqual(self.storage_snapshot(), before)

    def test_active_task_prevents_preparation(self) -> None:
        with self.app.state.database.session() as session:
            job = session.get(Job, self.export_id)
            job.status = "queued"
            session.commit()
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertFalse(self.candidate.exists())

    def test_archive_and_registered_source_mismatch_is_refused(self) -> None:
        self.results[0].write_bytes(b"different")
        before = self.storage_snapshot()
        with self.assertRaises(ValueError):
            self.prepare()
        self.assertEqual(self.storage_snapshot(), before)
        self.assertFalse(self.candidate.exists())

    def test_matching_archive_with_wrong_business_quantity_is_refused(self) -> None:
        book = load_workbook(self.results[0])
        book["交货导入"].cell(4, 4).value = 79
        book.save(self.results[0])
        book.close()
        self.rebuild_archive()
        before = self.storage_snapshot()
        with self.assertRaisesRegex(ValueError, "计算及审校结果"):
            self.prepare()
        self.assertEqual(self.storage_snapshot(), before)

    def test_changed_batch_or_candidate_prevents_publication(self) -> None:
        self.prepare()
        candidate = self.candidate / self.target.name
        original = candidate.read_bytes()
        candidate.write_bytes(original + b"changed")
        with self.assertRaises(ValueError):
            self.publish()
        candidate.write_bytes(original)
        with self.app.state.database.session() as session:
            batch = session.get(Batch, self.batch_id)
            batch.name = "changed"
            session.commit()
        before = self.storage_snapshot()
        with self.assertRaises(ValueError):
            self.publish()
        self.assertEqual(self.storage_snapshot(), before)

    def test_publish_race_preserves_concurrent_target_and_removes_temporary_file(
        self,
    ) -> None:
        self.prepare()
        link = os.link

        def concurrent_target(source: Path, target: Path) -> None:
            target.write_bytes(b"concurrent-result")
            link(source, target)

        with (
            patch(
                "scripts.merged_export_recovery.os.link", side_effect=concurrent_target
            ),
            self.assertRaises(FileExistsError),
        ):
            self.publish()
        self.assertEqual(self.target.read_bytes(), b"concurrent-result")
        self.assertFalse(list(self.target.parent.glob(".merged-recovery-*")))

    def test_copy_failure_leaves_no_target_or_temporary_file(self) -> None:
        self.prepare()
        before = self.storage_snapshot()
        with (
            patch(
                "scripts.merged_export_recovery.shutil.copyfileobj",
                side_effect=OSError("copy failed"),
            ),
            self.assertRaises(OSError),
        ):
            self.publish()
        self.assertEqual(self.storage_snapshot(), before)
        self.assertFalse(list(self.target.parent.glob(".merged-recovery-*")))

    def test_prepare_rejects_changed_state_and_removes_its_candidate(self) -> None:
        def write_and_change(*args, **kwargs) -> None:
            _write_merged_delivery(*args, **kwargs)
            with self.app.state.database.session() as session:
                batch = session.get(Batch, self.batch_id)
                batch.name = "changed-during-prepare"
                session.commit()

        with (
            patch(
                "scripts.merged_export_recovery._write_merged_delivery",
                side_effect=write_and_change,
            ),
            self.assertRaisesRegex(ValueError, "准备期间"),
        ):
            self.prepare()
        self.assertFalse(self.candidate.exists())
        self.assertFalse(self.target.exists())

    def test_wrong_pending_quantity_is_refused(self) -> None:
        book = load_workbook(self.results[1])
        book["待处理导入"].cell(3, 4).value = 34
        book.save(self.results[1])
        book.close()
        self.rebuild_archive()
        with self.assertRaisesRegex(ValueError, "待处理内容"):
            self.prepare()
        self.assertFalse(self.candidate.exists())

    def test_cli_refusal_does_not_leak_connection_credentials(self) -> None:
        result = subprocess.run(
            [
                sys.executable,
                "-B",
                "-m",
                "scripts.merged_export_recovery",
                "--database-url",
                "postgresql+invalid://user:secret-value@localhost/test",
                "--storage-root",
                str(self.storage_root),
                "--batch-id",
                str(self.batch_id),
                "--directory",
                str(self.candidate),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        self.assertEqual(result.returncode, 2)
        self.assertNotIn("secret-value", result.stdout + result.stderr)
        self.assertFalse(self.candidate.exists())
