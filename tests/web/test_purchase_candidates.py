from tests.support.web_api import WebApiCase
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd


from delivery_note.web.models import (
    InputVersion,
    PurchaseSyncJob,
    SelfOperatedInboundSyncJob,
)


class WebApiTests(WebApiCase):
    def test_purchase_sync_issues_can_be_previewed_by_an_operator(self):
        admin_headers = self.login("admin", "admin-pass")
        with self.app.state.database.session() as session:
            job = PurchaseSyncJob(
                status="succeeded",
                created_by=1,
                issues=[
                    {
                        "severity": "warning",
                        "message": "共享站点数据不能参与正常交货匹配",
                        "po_code": "PO-1001",
                        "sku": "SKU-A",
                        "source_site": "共享",
                        "supplier_code": "SUP-1",
                        "supplier_name": "供应商 A",
                        "warehouse": "水鞋-广州仓",
                        "quantity": 12,
                        "code": "shared_site",
                        "internal_note": "不导出的内部字段",
                    },
                    {
                        "severity": "error",
                        "message": "数量为零",
                        "po_code": "PO-1002",
                        "sku": "SKU-B",
                        "quantity": 0,
                        "code": "zero_quantity",
                    },
                    {
                        "severity": "warning",
                        "message": "小数数量待确认",
                        "po_code": "PO-1003",
                        "sku": "SKU-C",
                        "quantity": 0.5,
                    },
                ],
            )
            session.add(job)
            session.commit()
            job_id = job.id

        preview = self.client.get(
            f"/api/purchase-sync/{job_id}/issues",
            headers=admin_headers,
        )

        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()[0]["po_code"], "PO-1001")
        self.assertEqual(preview.json()[0]["warehouse"], "水鞋-广州仓")
        self.assertEqual(preview.json()[0]["quantity"], 12)
        download = self.client.get(
            f"/api/purchase-sync/{job_id}/issues/download",
            headers=admin_headers,
        )
        self.assertEqual(download.status_code, 200, download.text)
        issue_frame = pd.read_excel(BytesIO(download.content))
        self.assertEqual(issue_frame.loc[0, "目的仓"], "水鞋-广州仓")
        self.assertEqual(issue_frame.loc[0, "未交量"], 12)
        workbook = load_workbook(BytesIO(download.content))
        self.assertEqual(workbook.sheetnames, ["Sheet1"])
        self.assertEqual(
            list(workbook.active.values),
            [
                (
                    "级别",
                    "问题",
                    "采购单号",
                    "SKU",
                    "目的仓",
                    "未交量",
                    "接口站点",
                    "接口供应商编号",
                    "接口供应商名称",
                    "问题类型",
                ),
                (
                    "warning",
                    "共享站点数据不能参与正常交货匹配",
                    "PO-1001",
                    "SKU-A",
                    "水鞋-广州仓",
                    12,
                    "共享",
                    "SUP-1",
                    "供应商 A",
                    "shared_site",
                ),
                (
                    "error",
                    "数量为零",
                    "PO-1002",
                    "SKU-B",
                    None,
                    0,
                    None,
                    None,
                    None,
                    "zero_quantity",
                ),
                (
                    "warning",
                    "小数数量待确认",
                    "PO-1003",
                    "SKU-C",
                    None,
                    0.5,
                    None,
                    None,
                    None,
                    None,
                ),
            ],
        )
        self.assertEqual(
            download.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            download.headers["content-disposition"],
            f'attachment; filename="purchase_sync_issues_{job_id}.xlsx"',
        )
        workbook.close()
        operator = self.create_operator(admin_headers)
        operator_headers = self.login(operator["username"], "operator-pass")
        operator_preview = self.client.get(
            f"/api/purchase-sync/{job_id}/issues",
            headers=operator_headers,
        )
        self.assertEqual(operator_preview.status_code, 200, operator_preview.text)
        self.assertEqual(operator_preview.json()[0]["po_code"], "PO-1001")
        operator_download = self.client.get(
            f"/api/purchase-sync/{job_id}/issues/download",
            headers=operator_headers,
        )
        self.assertEqual(operator_download.status_code, 200, operator_download.text)
        pd.testing.assert_frame_equal(
            pd.read_excel(BytesIO(operator_download.content)), issue_frame
        )

    def test_purchase_sync_candidate_can_be_previewed_by_an_operator(self):
        admin_headers = self.login("admin", "admin-pass")
        candidate_path = self.root / "purchase-candidate.xlsx"
        workbook = load_workbook(BytesIO(self.workbook_bytes("purchase")))
        workbook.active.append(
            ["待交货", "KuangBiao", "SKU-B", "AMAZON:SEEKWAY:US", None, 20]
        )
        workbook.save(candidate_path)
        with self.app.state.database.session() as session:
            version = InputVersion(
                kind="purchase",
                name="采购候选版本",
                original_name="purchase-candidate.xlsx",
                storage_path=str(candidate_path),
                active=False,
                created_by=1,
            )
            session.add(version)
            session.flush()
            job = PurchaseSyncJob(
                status="succeeded",
                created_by=1,
                candidate_version_id=version.id,
            )
            session.add(job)
            session.commit()
            job_id = job.id

        operator = self.create_operator(admin_headers)
        operator_headers = self.login(operator["username"], "operator-pass")
        preview = self.client.get(
            f"/api/purchase-sync/{job_id}/preview?limit=100",
            headers=operator_headers,
        )

        self.assertEqual(preview.status_code, 200, preview.text)
        payload = preview.json()
        self.assertEqual(
            payload["columns"],
            ["单据状态", "供应商", "SKU", "平台站点", "目的仓", "未交量"],
        )
        self.assertEqual(payload["total"], 2)
        self.assertEqual(payload["rows"][0]["_row_number"], 1)
        self.assertEqual(payload["rows"][0]["SKU"], "SKU-A")
        self.assertEqual(payload["rows"][0]["目的仓"], "水鞋-广州仓")
        self.assertEqual(payload["rows"][0]["未交量"], 100)
        self.assertEqual(payload["rows"][1]["_row_number"], 2)
        self.assertEqual(payload["rows"][1]["SKU"], "SKU-B")
        self.assertIsNone(payload["rows"][1]["目的仓"])

        limited = self.client.get(
            f"/api/purchase-sync/{job_id}/preview?limit=1",
            headers=operator_headers,
        )
        self.assertEqual(limited.status_code, 200, limited.text)
        self.assertEqual(limited.json()["total"], 2)
        self.assertEqual(limited.json()["rows"], payload["rows"][:1])

    def test_sync_candidate_preview_error_statuses(self):
        headers = self.login("admin", "admin-pass")
        with self.app.state.database.session() as session:
            version = InputVersion(
                kind="purchase",
                name="无法读取的候选版本",
                original_name="missing.xlsx",
                storage_path=str(self.root / "missing.xlsx"),
                active=False,
                created_by=1,
            )
            session.add(version)
            session.flush()
            jobs = []
            for job_type in (PurchaseSyncJob, SelfOperatedInboundSyncJob):
                for candidate_version_id in (None, version.id):
                    job = job_type(
                        status="succeeded",
                        created_by=1,
                        candidate_version_id=candidate_version_id,
                    )
                    session.add(job)
                    session.flush()
                    jobs.append(job.id)
            session.commit()

        for prefix, job_ids in (
            ("purchase-sync", jobs[:2]),
            ("self-operated-inbound-sync", jobs[2:]),
        ):
            with self.subTest(prefix=prefix):
                missing_job = self.client.get(
                    f"/api/{prefix}/999/preview", headers=headers
                )
                no_candidate = self.client.get(
                    f"/api/{prefix}/{job_ids[0]}/preview", headers=headers
                )
                unreadable = self.client.get(
                    f"/api/{prefix}/{job_ids[1]}/preview", headers=headers
                )
                self.assertEqual(missing_job.status_code, 404)
                self.assertEqual(no_candidate.status_code, 409)
                self.assertEqual(unreadable.status_code, 409)
