from tests.support.web_api import WebApiCase
from io import BytesIO

from openpyxl import load_workbook

from delivery_note.processing.models import IMPORT_COLUMNS
from delivery_note.self_operated_inbound import INBOUND_TEMPLATE_COLUMNS


class WebApiTests(WebApiCase):
    def test_initial_state_includes_builtin_templates(self):
        admin_headers = self.login("admin", "admin-pass")
        versions = self.client.get(
            "/api/input-versions",
            headers=admin_headers,
        ).json()
        export_versions = [
            version for version in versions if version["kind"] == "template"
        ]
        inbound_versions = [
            version for version in versions if version["kind"] == "inbound_template"
        ]

        self.assertEqual(len(export_versions), 1)
        export_version = export_versions[0]
        self.assertEqual(
            export_version["name"],
            "系统内置交货导出模板",
        )
        self.assertEqual(
            export_version["original_name"],
            "交货导入模板.xlsx",
        )
        self.assertTrue(export_version["active"])

        export_download = self.client.get(
            f"/api/input-versions/{export_version['id']}/download",
            headers=admin_headers,
        )
        self.assertEqual(export_download.status_code, 200, export_download.text)
        export_workbook = load_workbook(
            BytesIO(export_download.content),
            read_only=True,
        )
        try:
            self.assertEqual(
                [
                    export_workbook.active.cell(row=2, column=column).value
                    for column in range(1, len(IMPORT_COLUMNS) + 1)
                ],
                IMPORT_COLUMNS,
            )
            self.assertEqual(export_workbook.active["A3"].value, "示例仓库")
        finally:
            export_workbook.close()

        self.assertEqual(len(inbound_versions), 1)
        version = inbound_versions[0]
        self.assertEqual(version["name"], "系统内置积加入库模板")
        self.assertEqual(version["original_name"], "积加批量入库模板.xlsx")
        self.assertTrue(version["active"])

        downloaded = self.client.get(
            f"/api/input-versions/{version['id']}/download",
            headers=admin_headers,
        )
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        workbook = load_workbook(BytesIO(downloaded.content), read_only=True)
        try:
            headers = [
                workbook.active.cell(row=1, column=column).value
                for column in range(1, len(INBOUND_TEMPLATE_COLUMNS) + 1)
            ]
            self.assertEqual(headers, INBOUND_TEMPLATE_COLUMNS)
            self.assertEqual(
                sum(
                    1 for sheet in workbook.worksheets if sheet.sheet_state == "hidden"
                ),
                70,
            )
        finally:
            workbook.close()

    def test_delivery_batch_uses_builtin_template_without_upload(self):
        admin_headers = self.login("admin", "admin-pass")
        for kind in ("purchase", "product", "supplier", "position"):
            response = self.client.post(
                f"/api/input-versions/{kind}",
                headers=admin_headers,
                data={"name": f"{kind}-v1", "activate": "true"},
                files={
                    "file": (
                        f"{kind}.xlsx",
                        BytesIO(self.workbook_bytes(kind)),
                    )
                },
            )
            self.assertEqual(response.status_code, 201, response.text)

        created = self.client.post(
            "/api/batches",
            headers=admin_headers,
            json={"name": "使用系统内置导出模板"},
        )

        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(
            created.json()["versions"]["template"]["name"],
            "系统内置交货导出模板",
        )
