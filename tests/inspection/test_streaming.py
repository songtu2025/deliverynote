from collections.abc import Iterator
from types import SimpleNamespace
from typing import Any
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from delivery_note.excel_io import (
    PRODUCT_COLUMNS,
    PURCHASE_COLUMNS,
)
from delivery_note.input_inspection import (
    inspect_input_version,
    inspect_input_version_with_preview,
    preview_input_version_page,
)
from tests.support.inspection import InputInspectionCase


class InputInspectionTests(InputInspectionCase):
    def test_combined_product_inspection_streams_only_required_columns(self) -> None:
        path = self.root / "product.xlsx"
        source_columns = [
            "其他字段",
            "锁仓MKSU",
            "SKU",
            "店铺/站点",
            "品类A",
        ]
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(source_columns)
        sheet.append(["忽略", "锁", "SKU-A", "SEEKWAY:US", "水鞋"])
        sheet.append(["忽略", "", "SKU-B", "SEEKWAY:CA", "配件"])
        sheet["A100"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        workbook.save(path)

        with patch(
            "delivery_note.inspection.workbooks.read_product_workbook",
            side_effect=AssertionError("产品预览不应全量读取 DataFrame"),
        ):
            result = inspect_input_version_with_preview(
                "product",
                path,
                offset=1,
                limit=1,
            )

        expected_columns = [
            column for column in source_columns if column in PRODUCT_COLUMNS
        ]
        self.assertEqual(result["summary"]["row_count"], 2)
        self.assertEqual(result["summary"]["columns"], expected_columns)
        self.assertEqual(result["preview"]["total"], 2)
        self.assertEqual(result["preview"]["columns"], expected_columns)
        self.assertEqual(
            result["preview"]["rows"],
            [
                {
                    "锁仓MKSU": None,
                    "SKU": "SKU-B",
                    "店铺/站点": "SEEKWAY:CA",
                    "品类A": "配件",
                }
            ],
        )

    def test_streaming_next_page_stops_at_page_end(self) -> None:
        consumed_rows: list[int] = []

        def workbook_rows() -> Iterator[tuple[Any, ...]]:
            yield tuple(PRODUCT_COLUMNS)
            for index in range(100):
                consumed_rows.append(index)
                yield (f"SKU-{index}", "SEEKWAY:US", "水鞋", "锁")

        workbook = SimpleNamespace(
            worksheets=[
                SimpleNamespace(
                    iter_rows=lambda **_kwargs: workbook_rows(),
                )
            ],
            close=lambda: None,
        )
        summary = {
            "kind": "product",
            "row_count": 100,
            "columns": PRODUCT_COLUMNS,
            "metrics": {},
            "issues": [],
        }

        with patch(
            "delivery_note.inspection.streaming.load_workbook",
            return_value=workbook,
        ):
            preview = preview_input_version_page(
                "product",
                self.root / "product.xlsx",
                offset=20,
                limit=10,
                summary=summary,
            )

        self.assertEqual(consumed_rows, list(range(30)))
        self.assertEqual(preview["total"], 100)
        self.assertEqual(preview["offset"], 20)
        self.assertEqual(preview["limit"], 10)
        self.assertEqual(
            [row["SKU"] for row in preview["rows"]],
            [f"SKU-{index}" for index in range(20, 30)],
        )

    def test_purchase_inspection_marks_shared_site_as_warning(self) -> None:
        path = self.root / "purchase.xlsx"
        frame = pd.DataFrame(
            [
                ["待交货", "接口供应商", "SKU-A", "共享", "广州仓", 12],
                [
                    "交货中",
                    "接口供应商",
                    "SKU-B",
                    "AMAZON:SEEKWAY:US",
                    "广州仓",
                    8,
                ],
            ],
            columns=PURCHASE_COLUMNS,
        )
        frame.to_excel(path, index=False)

        inspection = inspect_input_version("purchase", path)
        combined = inspect_input_version_with_preview(
            "purchase",
            path,
            offset=0,
            limit=20,
        )

        issue = inspection["issues"][0]
        self.assertEqual(issue["severity"], "warning")
        self.assertEqual(issue["code"], "shared_site")
        self.assertEqual(issue["row_numbers"], [2])
        self.assertEqual(combined["summary"]["issues"], [issue])
