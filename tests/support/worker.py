from typing import Any
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from delivery_note.processing.models import IMPORT_COLUMNS
from delivery_note.self_operated_inbound import INBOUND_TEMPLATE_COLUMNS
from tests.asgi_client import SyncASGIClient
from delivery_note.web.api import create_app


class WorkerCase(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.database_url = f"sqlite+pysqlite:///{self.root / 'worker.db'}"
        self.storage_root = self.root / "storage"
        self.app = create_app(
            database_url=self.database_url,
            storage_root=self.storage_root,
            bootstrap_admin=("admin", "admin-pass"),
        )
        self.client = SyncASGIClient(self.app)
        login = self.client.post(
            "/api/auth/login",
            json={"username": "admin", "password": "admin-pass"},
        )
        self.headers = {"Authorization": f"Bearer {login.json()['token']}"}
        self.inputs = self.create_master_inputs(self.root)
        self.upload_versions()

    def tearDown(self):
        if hasattr(self, "client"):
            self.client.close()
        if hasattr(self, "app"):
            self.app.state.database.dispose()
        if hasattr(self, "directory"):
            self.directory.cleanup()

    @staticmethod
    def create_master_inputs(root: Path) -> dict[str, Path]:
        purchase = root / "purchase.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(["单据状态", "供应商", "SKU", "平台站点", "目的仓", "未交量"])
        sheet.append(
            ["待交货", "KuangBiao", "SKU-A", "AMAZON:SEEKWAY:US", "水鞋-广州仓", 100]
        )
        workbook.save(purchase)

        product = root / "product.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(["SKU", "店铺/站点", "品类A", "锁仓MKSU"])
        sheet.append(["SKU-A", "SEEKWAY:US", "水鞋", "锁"])
        workbook.save(product)

        supplier = root / "supplier.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(["供应商编号", "供应商名称", "状态", "供应商别名"])
        sheet.append(["GYS-023", "KuangBiao", "启用", "瑞智雅|RIVBOS"])
        workbook.save(supplier)

        position = root / "position.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "MSKU_视图"
        sheet.append(
            ["店铺-站点", "积加SKU", "MSKU", "规模定位", "备货定位", "已下单可售天数"]
        )
        sheet.append(["SEEKWAY:US", "SKU-A", "MSKU-A", "短尾", "备货", 90])
        workbook.save(position)

        template = root / "template.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet["A1"] = "模板提示"
        sheet.merge_cells("A1:G1")
        sheet.append(IMPORT_COLUMNS)
        sheet.append(["示例仓", "示例供应商", "示例SKU", 1, "示例站点", "", ""])
        workbook.save(template)

        inbound_template = root / "inbound_template.xlsx"
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "批量入库"
        sheet.append(INBOUND_TEMPLATE_COLUMNS)
        sheet.append(["示例"] + [""] * (len(INBOUND_TEMPLATE_COLUMNS) - 1))
        workbook.save(inbound_template)
        return {
            "purchase": purchase,
            "product": product,
            "supplier": supplier,
            "position": position,
            "template": template,
            "inbound_template": inbound_template,
        }

    @staticmethod
    def create_delivery(path: Path, quantity: int) -> Path:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "明细"
        sheet.append([])
        sheet.append([])
        sheet.append([])
        sheet.append(["积加SKU", "数量", "站点"])
        sheet.append(["SKU-A", quantity, "US站"])
        workbook.save(path)
        return path

    @staticmethod
    def create_self_operated_delivery(path: Path, quantity: int) -> Path:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "明细"
        sheet.append([])
        sheet.append([])
        sheet.append([])
        sheet.append(["积加SKU", "实收数量", "站点", "交货单号"])
        sheet.append(["SKU-A", quantity, "US站", "LN2608179025"])
        workbook.save(path)
        return path

    @staticmethod
    def create_self_operated_inbound(path: Path) -> Path:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        source_columns = [
            column for column in INBOUND_TEMPLATE_COLUMNS if column != "最大可收货"
        ]
        sheet.append(source_columns + ["供应商"])
        for site, inbound_number in (
            ("AMAZON:RIVMOUNT:US", "IN-R"),
            ("AMAZON:SEEKWAY:US", "IN-S"),
        ):
            values: dict[str, Any] = {column: "" for column in INBOUND_TEMPLATE_COLUMNS}
            values.update(
                {
                    "入库单号": inbound_number,
                    "入库仓": "自营仓",
                    "SKU": "SKU-A",
                    "平台站点": site,
                    "关联采购单": f"PO-20260801-{site}",
                    "关联交货单/调拨单": "LN2608179025",
                    "应收货": 10,
                }
            )
            sheet.append([values[column] for column in source_columns] + ["KuangBiao"])
        workbook.save(path)
        return path

    def upload_versions(self):
        for kind, path in self.inputs.items():
            with path.open("rb") as upload:
                response = self.client.post(
                    f"/api/input-versions/{kind}",
                    headers=self.headers,
                    data={"name": f"{kind}-v1", "activate": "true"},
                    files={"file": (path.name, upload)},
                )
            self.assertEqual(response.status_code, 201, response.text)

    def create_batch(self, delivery_paths: list[Path]) -> tuple[int, int]:
        created = self.client.post(
            "/api/batches",
            headers=self.headers,
            json={"name": "Worker 集成测试"},
        )
        self.assertEqual(created.status_code, 201, created.text)
        batch_id = created.json()["id"]
        for path in delivery_paths:
            with path.open("rb") as upload:
                response = self.client.post(
                    f"/api/batches/{batch_id}/files",
                    headers=self.headers,
                    files={"file": (path.name, upload)},
                )
            self.assertEqual(response.status_code, 201, response.text)
        preflight = self.client.post(
            f"/api/batches/{batch_id}/preflight",
            headers=self.headers,
        )
        self.assertEqual(preflight.status_code, 200, preflight.text)
        job = self.client.post(
            f"/api/batches/{batch_id}/compute",
            headers=self.headers,
        )
        self.assertEqual(job.status_code, 202, job.text)
        return batch_id, job.json()["id"]
