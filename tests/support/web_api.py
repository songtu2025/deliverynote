from io import BytesIO
from typing import Any
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook
from sqlalchemy import event

from delivery_note.processing.models import IMPORT_COLUMNS
from delivery_note.inbound.models import (
    INBOUND_TEMPLATE_COLUMNS,
)
from tests.asgi_client import SyncASGIClient

from delivery_note.web.api import create_app


INPUT_KINDS = ("purchase", "product", "supplier", "position", "template")


class WebApiCase(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        root = Path(self.directory.name)
        self.root = root
        self.app = create_app(
            database_url=f"sqlite+pysqlite:///{root / 'test.db'}",
            storage_root=root / "storage",
            bootstrap_admin=("admin", "admin-pass"),
            session_cookie_secure=False,
        )
        self.client = SyncASGIClient(self.app)

    def tearDown(self):
        if hasattr(self, "client"):
            self.client.close()
        if hasattr(self, "app"):
            self.app.state.database.dispose()
        if hasattr(self, "directory"):
            self.directory.cleanup()

    def login(self, username: str, password: str) -> dict[str, str]:
        response = self.client.post(
            "/api/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return {"Authorization": f"Bearer {response.json()['token']}"}

    def create_operator(self, admin_headers: dict[str, str]) -> dict:
        response = self.client.post(
            "/api/users",
            headers=admin_headers,
            json={
                "username": "operator",
                "password": "operator-pass",
                "role": "operator",
            },
        )
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    @staticmethod
    def workbook_bytes(kind: str) -> bytes:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        if kind == "purchase":
            sheet.append(["单据状态", "供应商", "SKU", "平台站点", "目的仓", "未交量"])
            sheet.append(
                [
                    "待交货",
                    "KuangBiao",
                    "SKU-A",
                    "AMAZON:SEEKWAY:US",
                    "水鞋-广州仓",
                    100,
                ]
            )
        elif kind == "product":
            sheet.append(["SKU", "店铺/站点", "品类A", "锁仓MKSU"])
            sheet.append(["SKU-A", "SEEKWAY:US", "水鞋", "锁"])
        elif kind == "supplier":
            sheet.append(["供应商编号", "供应商名称", "状态"])
            sheet.append(["GYS-023", "KuangBiao", "启用"])
        elif kind == "position":
            sheet.title = "MSKU_视图"
            sheet.append(
                [
                    "店铺-站点",
                    "积加SKU",
                    "MSKU",
                    "规模定位",
                    "备货定位",
                    "已下单可售天数",
                ]
            )
            sheet.append(["SEEKWAY:US", "SKU-A", "MSKU-A", "短尾", "备货", 90])
        elif kind == "template":
            sheet["A1"] = "模板提示"
            sheet.merge_cells("A1:G1")
            sheet.append(IMPORT_COLUMNS)
            sheet.append(["示例仓", "示例供应商", "示例SKU", 1, "示例站点", "", ""])
        elif kind == "inbound_template":
            sheet.title = "批量入库"
            sheet.append(INBOUND_TEMPLATE_COLUMNS)
            sheet.append(["示例"] + [""] * (len(INBOUND_TEMPLATE_COLUMNS) - 1))
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def supplier_workbook_bytes(rows: list[list[str]]) -> bytes:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(["供应商编号", "供应商名称", "状态", "供应商别名"])
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def delivery_bytes(quantity: int = 40) -> bytes:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "明细"
        sheet.append([])
        sheet.append([])
        sheet.append([])
        sheet.append(["积加SKU", "数量", "站点"])
        sheet.append(["SKU-A", quantity, "US站"])
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def self_operated_delivery_bytes(quantity: int = 15) -> bytes:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.title = "明细"
        sheet.append([])
        sheet.append([])
        sheet.append([])
        sheet.append(["积加SKU", "实收数量", "站点", "交货单号"])
        sheet.append(["SKU-A", quantity, "US站", "LN2608179025"])
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def self_operated_inbound_bytes() -> bytes:
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        sheet.append(INBOUND_TEMPLATE_COLUMNS + ["供应商"])
        values: dict[str, Any] = {column: "" for column in INBOUND_TEMPLATE_COLUMNS}
        values.update(
            {
                "入库单号": "IN-1",
                "入库仓": "自营仓",
                "SKU": "SKU-A",
                "平台站点": "AMAZON:SEEKWAY:US",
                "关联采购单": "PO-20260801",
                "关联交货单/调拨单": "LN2608179025",
                "应收货": 10,
            }
        )
        sheet.append(
            [values[column] for column in INBOUND_TEMPLATE_COLUMNS] + ["KuangBiao"]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def upload_active_versions(self, headers: dict[str, str]) -> dict[str, int]:
        version_ids = {}
        for kind in INPUT_KINDS:
            response = self.client.post(
                f"/api/input-versions/{kind}",
                headers=headers,
                data={"name": f"{kind}-v1", "activate": "true"},
                files={"file": (f"{kind}.xlsx", BytesIO(self.workbook_bytes(kind)))},
            )
            self.assertEqual(response.status_code, 201, response.text)
            version_ids[kind] = response.json()["id"]
        return version_ids

    def get_with_query_count(
        self,
        path: str,
        headers: dict[str, str],
    ):
        statement_count = 0

        def count_statement(*_args):
            nonlocal statement_count
            statement_count += 1

        engine = self.app.state.database.engine
        event.listen(engine, "before_cursor_execute", count_statement)
        try:
            response = self.client.get(path, headers=headers)
        finally:
            event.remove(engine, "before_cursor_execute", count_statement)
        return response, statement_count
