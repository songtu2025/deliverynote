"""Excel 模板测试复用的最小工作簿。"""

from openpyxl import Workbook

from delivery_note.inbound.models import INBOUND_TEMPLATE_COLUMNS
from delivery_note.processing.models import IMPORT_COLUMNS


def make_import_template() -> Workbook:
    """构造保留提示、合并区域、表头和示例行的交货模板。"""

    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet["A1"] = "模板提示"
    sheet.merge_cells("A1:G1")
    sheet.append(IMPORT_COLUMNS)
    sheet.append(["示例仓", "示例供应商", "示例SKU", 1, "示例站点", "", ""])
    return workbook


def make_inbound_template() -> Workbook:
    """构造带原表头和第二行示例的积加入库模板。"""

    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "批量入库"
    sheet.append(INBOUND_TEMPLATE_COLUMNS)
    sheet.append(["示例"] * len(INBOUND_TEMPLATE_COLUMNS))
    return workbook
