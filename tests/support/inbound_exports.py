"""检查自营仓 A:T 内容、示例样式及单文件和压缩包契约。"""

from io import BytesIO
from pathlib import Path
from typing import Any
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from delivery_note.inbound.models import INBOUND_TEMPLATE_COLUMNS


def assert_inbound_exports(
    case: unittest.TestCase, record: dict[str, Any], *, receivable: int = 10
) -> None:
    batch = record["batch"]
    site = record["receipts"]["items"][0]["full_site"]
    multiple = len(batch["files"]) > 1
    expected_rows: list[tuple[Any, ...]] = []
    for source in batch["files"]:
        data = record["exports"][source["original_name"]]
        file_receipt = next(
            row for row in record["receipts"]["items"] if row["file_id"] == source["id"]
        )
        number = file_receipt["delivery_numbers"][0]
        remaining = (3 if number.endswith("9026") else 7) if multiple else receivable
        book = load_workbook(BytesIO(data), data_only=True)
        expected_rows.extend(
            tuple(cell.value for cell in row)
            for row in book.worksheets[0].iter_rows(min_row=2)
        )
        book.close()
        assert_inbound_workbook(
            case,
            data,
            site,
            (remaining, source["import_total"]),
            source["import_total"] > remaining,
            delivery_number=number,
        )
    if len(batch["files"]) == 1:
        case.assertEqual(record["download"], next(iter(record["exports"].values())))
        case.assertIsNone(record["merged"])
        return
    with ZipFile(BytesIO(record["download"])) as archive:
        expected = {
            Path(name).stem + "_积加入库.xlsx": data
            for name, data in record["exports"].items()
        }
        case.assertEqual(set(archive.namelist()), set(expected))
        for name, data in expected.items():
            case.assertEqual(archive.read(name), data)
    book = load_workbook(BytesIO(record["merged"]), data_only=True)
    sheet = book.worksheets[0]
    case.assertEqual([cell.value for cell in sheet[1]], INBOUND_TEMPLATE_COLUMNS)
    case.assertEqual(
        sorted(tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)),
        sorted(expected_rows),
    )
    case.assertEqual(
        sum(sheet.cell(row, 17).value for row in range(2, sheet.max_row + 1)), 15
    )
    book.close()


def assert_inbound_workbook(
    case: unittest.TestCase,
    data: bytes,
    site: str,
    quantities: tuple[int, int],
    overreceipt: bool,
    *,
    delivery_number: str = "LN2608179025",
) -> None:
    receivable, total = quantities
    workbook = load_workbook(BytesIO(data), data_only=True)
    sheet = workbook["批量入库"]
    case.assertEqual([cell.value for cell in sheet[1]], INBOUND_TEMPLATE_COLUMNS)
    case.assertEqual(sheet.max_row, 2)
    expected = {
        1: ("IN-R" if site == "AMAZON:RIVMOUNT:US" else "IN-S")
        + ("-second" if delivery_number.endswith("9026") else ""),
        3: "自营仓",
        5: "SKU-A",
        8: site,
        9: f"PO-20260801-{site}",
        10: delivery_number,
        12: receivable,
        17: total,
        18: "未分配库位",
        19: None,
        20: "规则允许超收：5" if overreceipt else None,
    }
    for column, value in expected.items():
        case.assertEqual(sheet.cell(2, column).value, value)
    for cell in sheet[2]:
        case.assertEqual(cell.font.name, "宋体")
        case.assertEqual(cell.font.sz, 10)
        case.assertEqual(cell.fill.fgColor.rgb, "00FFF2CC")
        case.assertEqual(cell.number_format, "0")
        case.assertEqual(cell.alignment.horizontal, "center")
        case.assertEqual(cell.alignment.vertical, "center")
        case.assertTrue(cell.alignment.wrap_text)
        for name in ("left", "right", "top", "bottom"):
            case.assertEqual(getattr(cell.border, name).style, "thin")
    case.assertEqual(sheet.row_dimensions[2].height, 26)
    workbook.close()
