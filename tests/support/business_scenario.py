"""通过真实 HTTP 请求复用交货测试数据，等待真实 Worker 完成任务。"""

from pathlib import Path
import time
from typing import Any, Literal, cast

from httpx2 import Client, Response
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from tests.support.delivery_exports import SPLIT_PARTS
from tests.support.worker import WorkerCase


class DeliveryScenario:
    def __init__(self, url: str, root: Path) -> None:
        self.root = root
        self.client = Client(base_url=url, timeout=10, trust_env=False)
        self.batch_id = 0

    def request(
        self, method: str, path: str, expected: int = 200, **kwargs: Any
    ) -> Response:
        response = self.client.request(method, path, **kwargs)
        if response.status_code != expected:
            raise AssertionError(
                f"{method} {path}: {response.status_code}, {response.text}"
            )
        return response

    def login(self) -> dict[str, Any]:
        result = self.request(
            "POST",
            "/api/auth/login",
            json={"username": "admin", "password": "admin-pass"},
        ).json()
        self.client.headers["Authorization"] = "Bearer " + result["token"]
        return cast(dict[str, Any], result["user"])

    def wait_job(
        self,
        identifier: int,
        seconds: int = 60,
        *,
        expected_status: Literal["succeeded", "failed"] = "succeeded",
    ) -> dict[str, Any]:
        deadline = time.monotonic() + seconds
        while True:
            job = self.request("GET", f"/api/jobs/{identifier}").json()
            if job["status"] == expected_status:
                return cast(dict[str, Any], job)
            if job["status"] in {"succeeded", "failed"} or time.monotonic() >= deadline:
                raise AssertionError(
                    f"Worker 任务未达到预期状态 {expected_status}：{job}"
                )
            time.sleep(0.1)

    def create_baseline(self) -> None:
        self.activate_inputs()
        self.create_delivery_batch()

    def activate_inputs(self) -> None:
        inputs = WorkerCase.create_master_inputs(self.root)
        workbook = load_workbook(inputs["template"])
        for cell in workbook.worksheets[0][3]:
            cell.font = Font(name="宋体", size=10, color="808080")
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        workbook.save(inputs["template"])
        workbook.close()
        workbook = load_workbook(inputs["inbound_template"])
        sheet = workbook.worksheets[0]
        sheet.row_dimensions[2].height = 26
        side = Side(style="thin", color="808080")
        for cell in sheet[2]:
            cell.font = Font(name="宋体", size=10, color="808080")
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.number_format = "0"
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(left=side, right=side, top=side, bottom=side)
        workbook.save(inputs["inbound_template"])
        workbook.close()
        for kind, path in inputs.items():
            with path.open("rb") as upload:
                self.request(
                    "POST",
                    f"/api/input-versions/{kind}",
                    201,
                    data={"name": f"{kind}-v1", "activate": "true"},
                    files={"file": (path.name, upload)},
                )

    def create_delivery_batch(self, *, wait: bool = True) -> int:
        self.batch_id = self.request(
            "POST", "/api/batches", 201, json={"name": "恢复业务基线"}
        ).json()["id"]
        for letter, boxes in (("A", 10), ("B", 20)):
            path = WorkerCase.create_delivery(
                self.root / f"260717-狂飙-{letter}交货单-发货{boxes}箱.xlsx", 80
            )
            with path.open("rb") as upload:
                self.request(
                    "POST",
                    f"/api/batches/{self.batch_id}/files",
                    201,
                    files={"file": (path.name, upload)},
                )
        self.request("POST", f"/api/batches/{self.batch_id}/preflight")
        job = self.request("POST", f"/api/batches/{self.batch_id}/compute", 202)
        identifier = int(job.json()["id"])
        if wait:
            self.wait_job(identifier)
        return identifier

    def batch(self) -> dict[str, Any]:
        return cast(
            dict[str, Any],
            self.request("GET", f"/api/batches/{self.batch_id}").json(),
        )

    def split_and_export(self) -> None:
        exceptions = self.request(
            "GET", f"/api/batches/{self.batch_id}/exceptions"
        ).json()
        if len(exceptions) != 1 or exceptions[0]["manual_quantity"] != 60:
            raise AssertionError(f"待处理数量与基线不一致：{exceptions}")
        self.request(
            "PUT",
            f"/api/exceptions/{exceptions[0]['id']}/split",
            json={"parts": SPLIT_PARTS},
        )
        self.export()

    def export(self, *, wait: bool = True) -> int:
        job = self.request("POST", f"/api/batches/{self.batch_id}/export", 202)
        identifier = int(job.json()["id"])
        if wait:
            self.wait_job(identifier)
        return identifier

    def downloads(self) -> tuple[bytes, bytes]:
        prefix = f"/api/batches/{self.batch_id}"
        merged = self.request("GET", prefix + "/download-merged")
        archive = self.request("GET", prefix + "/download")
        for response, filename in (
            (merged, f"batch-{self.batch_id}-merged.xlsx"),
            (archive, f"batch-{self.batch_id}.zip"),
        ):
            assert f'filename="{filename}"' in response.headers["content-disposition"]
        return merged.content, archive.content

    def snapshot(self) -> dict[str, Any]:
        versions = self.request("GET", "/api/input-versions").json()
        active = [version for version in versions if version["active"]]
        return {
            "batch": self.batch(),
            "versions": versions,
            "inspection": {
                version["kind"]: self.request(
                    "GET", f"/api/input-versions/{version['id']}/inspection"
                ).json()
                for version in active
            },
            "input_files": {
                version["kind"]: self.request(
                    "GET", f"/api/input-versions/{version['id']}/download"
                ).content
                for version in active
            },
            "exceptions": self.request(
                "GET", f"/api/batches/{self.batch_id}/exceptions"
            ).json(),
            "exports": self.downloads(),
        }
