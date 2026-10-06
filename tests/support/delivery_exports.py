"""共享交货端到端场景的 Excel 和 ZIP 业务断言。"""

from copy import copy
from io import BytesIO
import unittest
from typing import cast
from zipfile import ZipFile

from openpyxl import load_workbook

from delivery_note.processing.models import IMPORT_COLUMNS


SPLIT_PARTS = [
    {
        "quantity": 25,
        "destination": "水鞋-广州仓",
        "delivery_note": "超出采购未交量",
        "resolved": True,
    },
    {"quantity": 35, "delivery_note": "超出采购未交量", "resolved": False},
]


def assert_delivery_exports(
    case: unittest.TestCase,
    merged_payload: bytes,
    archive_payload: bytes,
    *,
    balance: int = 100,
    resolved: int = 25,
) -> None:
    merged_book = load_workbook(BytesIO(merged_payload), data_only=True)
    merged_import_sheet = merged_book["交货导入"]
    merged_pending_sheet = merged_book["待处理导入"]
    case.assertEqual(merged_import_sheet.cell(3, 1).value, "示例仓")
    case.assertEqual(
        [merged_import_sheet.cell(2, column).value for column in range(1, 8)],
        IMPORT_COLUMNS,
    )
    case.assertIn("A1:G1", merged_import_sheet.merged_cells)
    case.assertEqual(
        sum(
            cast(int, merged_import_sheet.cell(row, 4).value or 0)
            for row in range(4, merged_import_sheet.max_row + 1)
        ),
        balance + resolved,
    )
    case.assertEqual(
        sum(
            cast(int, merged_pending_sheet.cell(row, 4).value or 0)
            for row in range(3, merged_pending_sheet.max_row + 1)
        ),
        160 - balance - resolved,
    )
    merged_import_notes = [
        cast(str, merged_import_sheet.cell(row, 6).value)
        for row in range(4, merged_import_sheet.max_row + 1)
    ]
    case.assertTrue(merged_import_notes[0].endswith("-01-10箱"))
    case.assertTrue(all(note.endswith("-02-20箱") for note in merged_import_notes[1:]))
    case.assertTrue(
        cast(str, merged_pending_sheet.cell(3, 6).value).endswith("-02-20箱")
    )
    for column in range(1, 8):
        example = merged_import_sheet.cell(3, column)
        for row in range(4, merged_import_sheet.max_row + 1):
            cell = merged_import_sheet.cell(row, column)
            for attribute in ("font", "fill", "border", "alignment"):
                case.assertEqual(
                    copy(getattr(cell, attribute)), copy(getattr(example, attribute))
                )
            if column == 4:
                case.assertEqual(cell.number_format, "0")
    with ZipFile(BytesIO(archive_payload)) as archive:
        names = sorted(archive.namelist())
        case.assertEqual(
            names,
            [
                "260717-狂飙-A交货单-发货10箱_交货处理.xlsx",
                "260717-狂飙-B交货单-发货20箱_交货处理.xlsx",
            ],
        )
        second_book = load_workbook(BytesIO(archive.read(names[1])), data_only=True)
    import_sheet, pending_sheet = second_book["交货导入"], second_book["待处理导入"]
    import_total = sum(
        cast(int, import_sheet.cell(row, 4).value or 0)
        for row in range(4, import_sheet.max_row + 1)
    )
    pending_total = sum(
        cast(int, pending_sheet.cell(row, 4).value or 0)
        for row in range(3, pending_sheet.max_row + 1)
    )
    case.assertEqual(
        (import_total, pending_total),
        (balance - 80 + resolved, 160 - balance - resolved),
    )
    import_records = [
        [import_sheet.cell(row, column).value for column in range(1, 8)]
        for row in range(4, import_sheet.max_row + 1)
    ]
    case.assertEqual(len(import_records), 1)
    case.assertEqual(import_records[0][3], balance - 80 + resolved)
    case.assertEqual(import_records[0][6], f"超出采购未交量：{160 - balance}")
    merged_book.close()
    second_book.close()
