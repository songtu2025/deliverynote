"""三种导出工作簿的原有内容、样式和编辑权限场景。"""

from pathlib import Path
from tempfile import TemporaryDirectory
from typing import cast
import unittest

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from delivery_note.excel.styles import _StyledCell, _WorkbookStyles
from delivery_note.excel_io import (
    write_delivery_workbook,
    write_import_workbook,
    write_self_operated_inbound_workbook,
)
from delivery_note.inbound.models import INBOUND_TEMPLATE_COLUMNS
from delivery_note.processing.models import (
    BatchResult,
    EXCEPTION_COLUMNS,
    IMPORT_COLUMNS,
    PENDING_COLUMNS,
)
from tests.support.excel import make_import_template, make_inbound_template


class ExcelOutputTests(unittest.TestCase):
    def test_write_self_operated_inbound_workbook_replaces_template_rows(self) -> None:
        with TemporaryDirectory() as directory:
            template_path = Path(directory) / "批量入库模板.xlsx"
            output_path = Path(directory) / "批量入库结果.xlsx"
            workbook = make_inbound_template()
            sheet = workbook.worksheets[0]
            sheet["A2"].font = Font(bold=True)
            workbook.save(template_path)

            values: dict[str, str | int] = {
                column: "" for column in INBOUND_TEMPLATE_COLUMNS
            }
            values.update(
                {
                    "入库单号": "WV-1",
                    "SKU": "SKU-A",
                    "本次入库": 20,
                    "入库库位": "未分配库位",
                }
            )
            write_self_operated_inbound_workbook(
                template_path,
                output_path,
                pd.DataFrame([values]),
            )

            result = load_workbook(output_path)
            result_sheet = cast(Worksheet, result.active)
            self.assertEqual(result_sheet.max_row, 2)
            self.assertEqual(result_sheet["A2"].value, "WV-1")
            self.assertEqual(result_sheet["Q2"].value, 20)
            self.assertEqual(result_sheet["R2"].value, "未分配库位")
            self.assertTrue(result_sheet["A2"].font.bold)

    def test_write_import_workbook_preserves_template_header_and_data_style(
        self,
    ) -> None:
        with TemporaryDirectory() as directory:
            directory_path = Path(directory)
            template_path = directory_path / "template.xlsx"
            output_path = directory_path / "output.xlsx"
            workbook = make_import_template()
            sheet = workbook.worksheets[0]
            for cell in sheet[2]:
                cell.font = Font(color="FF0000", bold=True)
            for cell in sheet[3]:
                cell.font = Font(name="宋体", size=10, color="808080")
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            workbook.save(template_path)

            rows = pd.DataFrame(
                [
                    ["仓A", "KuangBiao", "SKU-A", 10, "AMAZON:SEEKWAY:US", "", ""],
                    ["仓B", "KuangBiao", "SKU-B", 20, "AMAZON:SEEKWAY:CA", "", ""],
                ],
                columns=IMPORT_COLUMNS,
            )
            write_import_workbook(template_path, output_path, rows)
            result = load_workbook(output_path)
            output_sheet = cast(Worksheet, result.active)

        self.assertEqual(output_sheet["A1"].value, "模板提示")
        self.assertEqual([cell.value for cell in output_sheet[2]], IMPORT_COLUMNS)
        self.assertEqual(output_sheet["A3"].value, "仓A")
        self.assertEqual(output_sheet["A4"].value, "仓B")
        self.assertEqual(output_sheet["D3"].value, 10)
        self.assertEqual(output_sheet["D4"].value, 20)
        self.assertEqual(
            cast(_StyledCell, output_sheet["A3"])._style,
            cast(_StyledCell, output_sheet["A4"])._style,
        )
        self.assertNotEqual(output_sheet["A3"].value, "示例仓")

    def test_write_delivery_workbook_contains_import_details_and_editable_pending_rows(
        self,
    ) -> None:
        result = BatchResult(
            import_rows=pd.DataFrame(
                [["仓A", "KuangBiao", "SKU-A", 80, "AMAZON:SEEKWAY:US", "", ""]],
                columns=IMPORT_COLUMNS,
            ),
            exception_rows=pd.DataFrame(
                [
                    [
                        "SKU-A",
                        "US",
                        "AMAZON:SEEKWAY:US",
                        "水鞋-广州仓",
                        120,
                        80,
                        40,
                        "超出采购未交量",
                    ]
                ],
                columns=EXCEPTION_COLUMNS,
            ),
            delivery_total=120,
            import_total=80,
            manual_total=40,
        )

        with TemporaryDirectory() as directory:
            template_path = Path(directory) / "template.xlsx"
            output_path = Path(directory) / "delivery.xlsx"
            template_book = make_import_template()
            template_sheet = template_book.worksheets[0]
            for cell in template_sheet[3]:
                cell.font = Font(name="宋体", size=10, color="808080")
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            template_sheet.row_dimensions[3].height = 24
            template_sheet.protection.sheet = True
            template_book.save(template_path)
            pending_rows = pd.DataFrame(
                [
                    [
                        "水鞋-广州仓",
                        "KuangBiao",
                        "SKU-A",
                        40,
                        "AMAZON:SEEKWAY:US",
                        "狂飙交货单",
                        "超出采购未交量：40",
                        '{"MSKU-B":"短尾","MSKU-A":"长尾"}',
                        '{"MSKU-B":"备货","MSKU-A":"不备货"}',
                    ]
                ],
                columns=PENDING_COLUMNS,
            )
            write_delivery_workbook(
                template_path,
                output_path,
                result,
                result.import_rows,
                pending_rows,
            )
            workbook = load_workbook(output_path, data_only=True)

        self.assertEqual(workbook.sheetnames, ["交货导入", "待处理导入"])
        self.assertNotIn("运行汇总", workbook.sheetnames)
        self.assertNotIn("异常明细", workbook.sheetnames)
        for sheet_name in ("交货导入", "待处理导入"):
            protection = workbook[sheet_name].protection
            self.assertTrue(protection.sheet)
            self.assertFalse(protection.formatCells)
            self.assertFalse(protection.formatColumns)
            self.assertFalse(protection.formatRows)
            self.assertFalse(protection.selectLockedCells)
        self._assert_delivery_sheet(workbook)
        self._assert_pending_sheet(workbook)

    def _assert_delivery_sheet(self, workbook: Workbook) -> None:
        """核对示例保留、正式数据和默认编辑权限。"""

        self.assertEqual(workbook["交货导入"]["A1"].value, "模板提示")
        self.assertEqual(workbook["交货导入"]["A3"].value, "示例仓")
        self.assertEqual(workbook["交货导入"]["A4"].value, "仓A")
        self.assertEqual(workbook["交货导入"]["D4"].value, 80)
        self.assertEqual(workbook["交货导入"]["A3"].font.name, "宋体")
        self.assertEqual(workbook["交货导入"]["A3"].font.size, 10)
        self.assertEqual(workbook["交货导入"]["A3"].font.color.rgb, "00808080")
        self.assertEqual(workbook["交货导入"]["A3"].fill.fill_type, "solid")
        self.assertEqual(workbook["交货导入"]["A3"].fill.fgColor.rgb, "00FFF2CC")
        self.assertEqual(workbook["交货导入"].row_dimensions[3].height, 24)
        self.assertTrue(workbook["交货导入"].protection.sheet)
        self.assertTrue(workbook["交货导入"]["A1"].protection.locked)
        self.assertTrue(workbook["交货导入"]["A2"].protection.locked)
        self.assertFalse(workbook["交货导入"]["A3"].protection.locked)
        self.assertFalse(workbook["交货导入"]["G4"].protection.locked)
        workbook_styles = cast(_WorkbookStyles, workbook)
        default_protection_id = workbook_styles._cell_styles[0].protectionId
        self.assertFalse(workbook_styles._protections[default_protection_id].locked)

    def _assert_pending_sheet(self, workbook: Workbook) -> None:
        """核对待处理数量、定位内容和表头保护。"""

        self.assertEqual(workbook["待处理导入"]["A1"].value, "模板提示")
        self.assertEqual(
            [cell.value for cell in workbook["待处理导入"][2]], PENDING_COLUMNS
        )
        self.assertEqual(workbook["待处理导入"]["A3"].value, "水鞋-广州仓")
        self.assertEqual(workbook["待处理导入"]["F3"].value, "狂飙交货单")
        self.assertTrue(workbook["待处理导入"].protection.sheet)
        self.assertTrue(workbook["待处理导入"]["A1"].protection.locked)
        self.assertTrue(workbook["待处理导入"]["A2"].protection.locked)
        self.assertFalse(workbook["待处理导入"]["A3"].protection.locked)
        self.assertFalse(workbook["待处理导入"]["G3"].protection.locked)

        self.assertEqual(
            workbook["待处理导入"]["H3"].value,
            '{"MSKU-B":"短尾","MSKU-A":"长尾"}',
        )
        self.assertIsNone(workbook["待处理导入"]["J2"].value)
        self.assertIsNone(workbook["待处理导入"]["J3"].value)
        self.assertTrue(workbook["待处理导入"]["H2"].protection.locked)
        self.assertFalse(workbook["待处理导入"]["H3"].protection.locked)
        self.assertTrue(workbook["待处理导入"]["H3"].alignment.wrap_text)
        self.assertEqual(workbook["待处理导入"].row_dimensions[3].height, 60)
        self.assertEqual(workbook["待处理导入"].column_dimensions["G"].width, 35)
