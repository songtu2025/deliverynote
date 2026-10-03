"""交货、供应商和定位工作簿的原有读取场景。"""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from openpyxl import Workbook

from delivery_note.excel_io import (
    read_delivery_workbook,
    read_position_workbook,
    read_supplier_workbook,
)
from delivery_note.processing.models import POSITION_SOURCE_COLUMNS


class ExcelInputTests(unittest.TestCase):
    def test_read_delivery_workbook_uses_detail_even_with_summary(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "明细交货单.xlsx"
            workbook = Workbook()
            sheet = workbook.worksheets[0]
            sheet.title = "明细"
            sheet.append([])
            sheet.append([])
            sheet.append([])
            sheet.append(["积加SKU", "数量", "站点"])
            sheet.append(["SKU-A", 5, "CA站"])
            sheet.append(["SKU-A", 10, "US站"])
            sheet.append(["SKU-A", 2, "CA站"])
            sheet.append(["合计", None, None])
            sheet.append(["次品问题描述", None, None])
            summary = workbook.create_sheet("汇总")
            summary.append(["SKU", "US站"])
            summary.append(["SKU-A", 999])
            workbook.save(path)

            result = read_delivery_workbook(path)

        self.assertEqual(
            result.to_dict("records"),
            [
                {"SKU": "SKU-A", "原始站点": "CA", "交货量": 7},
                {"SKU": "SKU-A", "原始站点": "US", "交货量": 10},
            ],
        )

    def test_read_delivery_workbook_requires_detail_sheet(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "只有汇总的交货单.xlsx"
            workbook = Workbook()
            sheet = workbook.worksheets[0]
            sheet.title = "汇总"
            sheet.append([])
            sheet.append(["SKU", "US站"])
            sheet.append(["SKU-A", 10])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "明细"):
                read_delivery_workbook(path)

    def test_read_delivery_workbook_rejects_missing_detail_columns(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "缺少数量的交货单.xlsx"
            workbook = Workbook()
            sheet = workbook.worksheets[0]
            sheet.title = "明细"
            sheet.append([])
            sheet.append([])
            sheet.append([])
            sheet.append(["积加SKU", "站点"])
            sheet.append(["SKU-A", "US站"])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "数量"):
                read_delivery_workbook(path)

    def test_read_added_supplier_xls(self) -> None:
        source = Path(__file__).parent.parent / "fixtures" / "supplier_minimal.xls"
        result = read_supplier_workbook(source)
        zhangdun = result[result["供应商名称"].eq("Zhangdun")].iloc[0]

        self.assertEqual(zhangdun["供应商编号"], "GYS-027")
        self.assertEqual(zhangdun["状态"], "启用")
        self.assertIn("供应商别名", result.columns)
        self.assertEqual(zhangdun["供应商别名"], "")

    def test_read_supplier_workbook_keeps_optional_alias_column(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "supplier.xlsx"
            pd.DataFrame(
                [["STGYS001", "RUIZY", "启用", "瑞智雅|RIVBOS"]],
                columns=["供应商编号", "供应商名称", "状态", "供应商别名"],
            ).to_excel(path, index=False)

            result = read_supplier_workbook(path)

        self.assertEqual(result.iloc[0]["供应商别名"], "瑞智雅|RIVBOS")

    def test_read_position_workbook_ignores_ordered_days(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "position.xlsx"
            pd.DataFrame(
                [["SEEKWAY:US", "SKU-A", "MSKU-A", "短尾", "备货", 90]],
                columns=[*POSITION_SOURCE_COLUMNS, "已下单可售天数"],
            ).to_excel(path, sheet_name="MSKU_视图", index=False)

            result = read_position_workbook(path)

        self.assertEqual(result.columns.tolist(), POSITION_SOURCE_COLUMNS)
        self.assertNotIn("已下单可售天数", result.columns)
