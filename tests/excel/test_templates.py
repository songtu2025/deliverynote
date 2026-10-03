"""模板表头、示例行和活动工作表选择的契约测试。"""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl.styles import Font

from delivery_note.excel_io import (
    validate_self_operated_template_workbook,
    validate_template_workbook,
)
from tests.support.excel import make_import_template, make_inbound_template


class TemplateTests(unittest.TestCase):
    def test_import_template_requires_original_headers(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.xlsx"
            workbook = make_import_template()
            workbook.save(path)
            validate_template_workbook(path)
            workbook.worksheets[0]["A2"] = "错误字段"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "官方模板表头与预期字段不一致"):
                validate_template_workbook(path)

    def test_import_template_requires_example_value_or_style(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.xlsx"
            workbook = make_import_template()
            sheet = workbook.worksheets[0]
            for cell in sheet[3]:
                cell.value = ""
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "官方模板缺少第 3 行示例格式"):
                validate_template_workbook(path)
            sheet["A3"].font = Font(bold=True)
            workbook.save(path)
            validate_template_workbook(path)

    def test_import_template_rejects_missing_and_sparse_empty_example(self) -> None:
        for mode in ("missing", "sparse"):
            with self.subTest(mode=mode), TemporaryDirectory() as directory:
                path = Path(directory) / "template.xlsx"
                workbook = make_import_template()
                sheet = workbook.worksheets[0]
                sheet.delete_rows(3)
                if mode == "sparse":
                    sheet["A4"] = "示例行之外的内容"
                workbook.save(path)
                with self.assertRaisesRegex(ValueError, "官方模板缺少第 3 行示例格式"):
                    validate_template_workbook(path)

    def test_import_template_accepts_value_or_style_in_any_column(self) -> None:
        for column in range(1, 8):
            for mode in ("value", "style"):
                with (
                    self.subTest(column=column, mode=mode),
                    TemporaryDirectory() as directory,
                ):
                    path = Path(directory) / "template.xlsx"
                    workbook = make_import_template()
                    sheet = workbook.worksheets[0]
                    for cell in sheet[3]:
                        cell.value = None
                    if mode == "value":
                        sheet.cell(row=3, column=column).value = "有效示例"
                    else:
                        sheet.cell(row=3, column=column).font = Font(bold=True)
                    workbook.save(path)
                    validate_template_workbook(path)

    def test_inbound_template_requires_headers_and_example_row(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.xlsx"
            workbook = make_inbound_template()
            workbook.save(path)
            validate_self_operated_template_workbook(path)
            workbook.worksheets[0]["A1"] = "错误字段"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "积加入库模板表头与预期字段不一致"):
                validate_self_operated_template_workbook(path)
            workbook = make_inbound_template()
            workbook.worksheets[0].delete_rows(2)
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "积加入库模板缺少第 2 行示例格式"):
                validate_self_operated_template_workbook(path)

    def test_validation_uses_active_sheet(self) -> None:
        for factory, validator in (
            (make_import_template, validate_template_workbook),
            (make_inbound_template, validate_self_operated_template_workbook),
        ):
            with (
                self.subTest(factory=factory.__name__),
                TemporaryDirectory() as directory,
            ):
                path = Path(directory) / "template.xlsx"
                workbook = factory()
                workbook.create_sheet("辅助工作表", 0)
                workbook.active = 1
                workbook.save(path)
                validator(path)
