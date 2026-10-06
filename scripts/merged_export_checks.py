"""按业务内容核验旧分文件导出，兼容历史数据从第三行开始的布局。"""

from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd

from delivery_note.excel.styles import _excel_value
from delivery_note.excel.templates import _validate_template_headers
from delivery_note.processing.models import IMPORT_COLUMNS, PENDING_COLUMNS
from delivery_note.workers.export_rows import _consolidate_import_rows


def _values(frame: pd.DataFrame) -> list[list[Any]]:
    return [
        [_excel_value(value) if value != "" else None for value in row]
        for row in frame.itertuples(index=False, name=None)
    ]


def _body(sheet: Worksheet, start: int, columns: list[str]) -> pd.DataFrame:
    _validate_template_headers(
        sheet, columns, row=2, message="历史导出字段与锁定结果不一致"
    )
    rows = list(sheet.iter_rows(min_row=start, max_col=len(columns), values_only=True))
    return pd.DataFrame(
        [row for row in rows if any(value is not None for value in row)],
        columns=columns,
    )


def check_source_workbook(
    payload: bytes,
    import_rows: pd.DataFrame,
    pending_rows: pd.DataFrame,
    template: Path,
) -> dict[str, int]:
    """比较数量、业务身份、备注和定位；旧样式差异单独报告。"""
    book = load_workbook(BytesIO(payload), data_only=False)
    template_book = load_workbook(template)
    try:
        example = cast(Worksheet, template_book.active)
        imported = book["交货导入"]
        sample = [example.cell(3, column).value for column in range(1, 8)]
        preserves_sample = [
            imported.cell(3, column).value for column in range(1, 8)
        ] == sample
        start = 4 if preserves_sample else 3
        original = _consolidate_import_rows(_body(imported, start, IMPORT_COLUMNS))
        if _values(original) != _values(import_rows[IMPORT_COLUMNS]):
            raise ValueError("历史导入内容与已保存的计算及审校结果不一致")
        pending = _body(book["待处理导入"], 3, PENDING_COLUMNS)
        if _values(pending) != _values(pending_rows[PENDING_COLUMNS]):
            raise ValueError("历史待处理内容与已保存的审校结果不一致")
        differences = 0
        for row in imported.iter_rows(min_row=start, max_col=7):
            for column, cell in enumerate(row, 1):
                for attribute in ("font", "fill", "border", "alignment"):
                    differences += copy(getattr(cell, attribute)) != copy(
                        getattr(example.cell(3, column), attribute)
                    )
        return {
            "legacy_layout": int(not preserves_sample),
            "style_differences": differences,
        }
    finally:
        book.close()
        template_book.close()
