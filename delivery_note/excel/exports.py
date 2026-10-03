"""沿用官方模板生成三种实际使用的导出工作簿。"""

from copy import copy
from pathlib import Path
from typing import cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..inbound.models import INBOUND_TEMPLATE_COLUMNS
from ..processing.models import IMPORT_COLUMNS, BatchResult
from .styles import (
    _StyledCell,
    _excel_value,
    _populate_pending_position_columns,
    _protect_header_only,
    _set_default_cells_unlocked,
)
from .templates import _populate_import_sheet, _validate_template_headers


def write_import_workbook(
    template_path: Path,
    output_path: Path,
    import_rows: pd.DataFrame,
) -> None:
    """复制官方模板，在第 3 行起写入可导入数据。"""
    workbook = load_workbook(template_path)
    _populate_import_sheet(cast(Worksheet, workbook.active), import_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_self_operated_inbound_workbook(
    template_path: Path,
    output_path: Path,
    allocation_rows: pd.DataFrame,
) -> None:
    """复制积加入库模板，从第 2 行写入本次可入库数据。"""
    missing = [
        column
        for column in INBOUND_TEMPLATE_COLUMNS
        if column not in allocation_rows.columns
    ]
    if missing:
        raise ValueError(f"积加入库数据缺少字段：{', '.join(missing)}")

    workbook = load_workbook(template_path)
    sheet = cast(Worksheet, workbook.active)
    _validate_template_headers(
        sheet,
        INBOUND_TEMPLATE_COLUMNS,
        row=1,
        message="积加入库模板表头与预期字段不一致",
    )
    if sheet.max_row < 2:
        raise ValueError("积加入库模板缺少第 2 行示例格式")

    styles = [
        copy(cast(_StyledCell, sheet.cell(row=2, column=column))._style)
        for column in range(1, len(INBOUND_TEMPLATE_COLUMNS) + 1)
    ]
    row_height = sheet.row_dimensions[2].height
    sheet.delete_rows(2, sheet.max_row - 1)
    export_rows = allocation_rows[INBOUND_TEMPLATE_COLUMNS]
    for row_offset, values in enumerate(
        export_rows.itertuples(index=False, name=None),
        start=2,
    ):
        sheet.row_dimensions[row_offset].height = row_height
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_offset, column=column)
            cast(_StyledCell, cell)._style = copy(styles[column - 1])
            cell.value = _excel_value(value)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_delivery_workbook(
    template_path: Path,
    output_path: Path,
    result: BatchResult,
    import_rows: pd.DataFrame,
    pending_rows: pd.DataFrame,
) -> None:
    """输出交货导入和可编辑的待处理数据，保持表头保护。"""
    workbook = load_workbook(template_path)
    import_sheet = cast(Worksheet, workbook.active)
    import_sheet.title = "交货导入"
    pending_sheet = workbook.copy_worksheet(import_sheet)
    pending_sheet.title = "待处理导入"
    pending_sheet.protection = copy(import_sheet.protection)

    _set_default_cells_unlocked(workbook)
    _populate_import_sheet(
        import_sheet,
        import_rows,
        preserve_example_row=True,
    )
    _populate_import_sheet(pending_sheet, pending_rows[IMPORT_COLUMNS])
    _populate_pending_position_columns(pending_sheet, pending_rows)
    _protect_header_only(import_sheet)
    _protect_header_only(pending_sheet)
    workbook.active = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
