"""交货和自营入库模板校验。"""

from pathlib import Path
from typing import cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..inbound.models import INBOUND_TEMPLATE_COLUMNS
from ..processing.models import IMPORT_COLUMNS


def _validate_template_headers(
    sheet: Worksheet,
    columns: list[str],
    *,
    row: int,
    message: str,
) -> None:
    """按模板原顺序检查表头，保持各模板的错误提示。"""

    headers = [
        sheet.cell(row=row, column=index).value for index in range(1, len(columns) + 1)
    ]
    if headers != columns:
        raise ValueError(message)


def validate_template_workbook(path: Path) -> None:
    """只读校验官方模板的 A:G 表头和示例格式行。"""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = cast(Worksheet, workbook.active)
        _validate_template_headers(
            sheet, IMPORT_COLUMNS, row=2, message="官方模板表头与预期字段不一致"
        )
        example_cells = [sheet.cell(row=3, column=column) for column in range(1, 8)]
        if not any(cell.value is not None or cell.has_style for cell in example_cells):
            raise ValueError("官方模板缺少第 3 行示例格式")
    finally:
        workbook.close()


def validate_self_operated_template_workbook(path: Path) -> None:
    """只读校验积加入库模板表头和样式示例行。"""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = cast(Worksheet, workbook.active)
        _validate_template_headers(
            sheet,
            INBOUND_TEMPLATE_COLUMNS,
            row=1,
            message="积加入库模板表头与预期字段不一致",
        )
        if sheet.max_row < 2:
            raise ValueError("积加入库模板缺少第 2 行示例格式")
    finally:
        workbook.close()
