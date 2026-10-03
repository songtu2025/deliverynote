"""产品与采购工作簿的流式摘要和分页读取。"""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..excel_io import PRODUCT_COLUMNS, PURCHASE_COLUMNS
from . import frames

_STREAMING_COLUMNS = {
    "product": PRODUCT_COLUMNS,
    "purchase": PURCHASE_COLUMNS,
}


def _stream_xlsx_inspection(
    kind: str,
    path: Path,
    offset: int,
    limit: int,
) -> dict[str, Any]:
    """流式读取预览行，同时计算有效数据总行数。"""

    expected_columns = _STREAMING_COLUMNS[kind]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        values = sheet.iter_rows(values_only=True)
        selected_columns = _stream_selected_columns(
            expected_columns,
            next(values, ()),
        )

        page_rows: list[tuple[int, dict[str, Any]]] = []
        shared_site_rows: list[int] = []
        site_column_index = next(
            (index for index, name in selected_columns if name == "平台站点"),
            None,
        )
        last_data_offset = -1
        for row_offset, row in enumerate(values):
            selected_values = _selected_values(selected_columns, row)
            if any(value not in (None, "") for value in selected_values):
                last_data_offset = row_offset
            if (
                kind == "purchase"
                and site_column_index is not None
                and site_column_index < len(row)
                and str(row[site_column_index] or "").strip() == "共享"
            ):
                shared_site_rows.append(row_offset + 2)
            if offset <= row_offset < offset + limit:
                page_rows.append(
                    (
                        row_offset,
                        _preview_row(selected_columns, selected_values),
                    )
                )

        total = last_data_offset + 1
        columns = [name for _index, name in selected_columns]
        preview: dict[str, Any] = {
            "kind": kind,
            "columns": columns,
            "rows": [row for row_offset, row in page_rows if row_offset < total],
            "total": total,
            "offset": offset,
            "limit": limit,
        }
        summary = {
            "kind": kind,
            "row_count": total,
            "columns": columns,
            "metrics": {},
            "issues": (frames.shared_site_warning(shared_site_rows)),
        }
        return {"summary": summary, "preview": preview}
    finally:
        workbook.close()


def _stream_selected_columns(
    expected_columns: list[str],
    header: tuple[Any, ...],
) -> list[tuple[int, str]]:
    """按源列顺序选取必要字段，并报告缺失字段。"""

    header_names = ["" if value is None else str(value) for value in header]
    expected_set = set(expected_columns)
    selected_columns = [
        (index, name) for index, name in enumerate(header_names) if name in expected_set
    ]
    found = {name for _index, name in selected_columns}
    missing = [column for column in expected_columns if column not in found]
    if missing:
        raise ValueError(f"缺少必要字段：{', '.join(missing)}")
    return selected_columns


def _stream_xlsx_preview(
    kind: str,
    path: Path,
    offset: int,
    limit: int,
    summary: dict[str, Any],
) -> dict[str, Any]:
    """使用已知总行数读取一页，并在页末停止流式扫描。"""

    total = int(summary["row_count"])
    columns = list(summary["columns"])
    preview: dict[str, Any] = {
        "kind": kind,
        "columns": columns,
        "rows": [],
        "total": total,
        "offset": offset,
        "limit": limit,
    }
    page_end = min(offset + limit, total)
    if offset >= page_end:
        return preview

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        values = sheet.iter_rows(values_only=True)
        selected_columns = _stream_selected_columns(
            _STREAMING_COLUMNS[kind],
            next(values, ()),
        )
        for row_offset, row in enumerate(values):
            if row_offset < offset:
                continue
            selected_values = _selected_values(selected_columns, row)
            preview["rows"].append(_preview_row(selected_columns, selected_values))
            if row_offset + 1 >= page_end:
                break
        return preview
    finally:
        workbook.close()


def _selected_values(
    selected_columns: list[tuple[int, str]], row: tuple[Any, ...]
) -> list[Any]:
    """从流式行中提取必要字段，保留原始空值判断。"""

    return [
        row[index] if index < len(row) else None for index, _name in selected_columns
    ]


def _preview_row(
    selected_columns: list[tuple[int, str]], selected_values: list[Any]
) -> dict[str, Any]:
    """共用摘要扫描和分页读取的预览字段序列化。"""

    return {
        name: frames._stream_json_safe(value)
        for (_index, name), value in zip(selected_columns, selected_values)
    }
